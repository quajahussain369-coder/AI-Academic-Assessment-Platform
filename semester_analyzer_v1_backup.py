import os
import pandas as pd
import matplotlib.pyplot as plt
import numpy as np

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, Reference

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image
)


# ============================================================
# 1. GRADE CALCULATION
# ============================================================

def calculate_grade(mark):
    """
    Example grading system.
    Change these ranges according to your college rules.
    """

    if mark >= 90:
        return "A+"
    elif mark >= 80:
        return "A"
    elif mark >= 70:
        return "B+"
    elif mark >= 60:
        return "B"
    elif mark >= 50:
        return "C"
    elif mark >= 40:
        return "D"
    else:
        return "F"


# ============================================================
# 2. GET STUDENT INFORMATION
# ============================================================

print("=" * 60)
print("        FIRST SEMESTER ECE PERFORMANCE ANALYZER")
print("=" * 60)

student_name = input("\nEnter student name: ")

semester = input("Enter semester (example: First Semester): ")

print("\nEnter the number of subjects.")
number_of_subjects = int(input("Number of subjects: "))


# ============================================================
# 3. ENTER SUBJECT DATA
# ============================================================

subjects = []
marks = []

print("\n" + "-" * 60)
print("ENTER SUBJECT DETAILS")
print("-" * 60)

for i in range(number_of_subjects):

    print(f"\nSubject {i + 1}")

    subject = input("Subject name: ")

    while True:
        try:
            mark = float(input("Marks out of 100: "))

            if 0 <= mark <= 100:
                break
            else:
                print("Please enter marks between 0 and 100.")

        except ValueError:
            print("Please enter a valid number.")

    subjects.append(subject)
    marks.append(mark)


# ============================================================
# 4. PROCESS THE DATA
# ============================================================

grades = [calculate_grade(mark) for mark in marks]

percentages = marks.copy()

total_marks = sum(marks)
maximum_marks = number_of_subjects * 100

overall_percentage = (total_marks / maximum_marks) * 100


# Create DataFrame

data = {
    "S.No.": range(1, number_of_subjects + 1),
    "Subject": subjects,
    "Marks": marks,
    "Percentage": percentages,
    "Grade": grades
}

df = pd.DataFrame(data)


# ============================================================
# 5. DISPLAY RESULTS
# ============================================================

print("\n")
print("=" * 75)
print("                 SEMESTER RESULT")
print("=" * 75)

print(df.to_string(index=False))

print("\n" + "-" * 75)
print(f"Total Marks       : {total_marks:.0f}/{maximum_marks}")
print(f"Overall Percentage: {overall_percentage:.2f}%")
print("-" * 75)


# ============================================================
# 6. CREATE OUTPUT FOLDER
# ============================================================

output_folder = "Semester_Result"

if not os.path.exists(output_folder):
    os.makedirs(output_folder)


# ============================================================
# 7. BAR CHART
# ============================================================

plt.figure(figsize=(10, 6))

bars = plt.bar(subjects, marks)

plt.title("First Semester ECE - Subject Performance")
plt.xlabel("Subjects")
plt.ylabel("Marks (%)")
plt.ylim(0, 100)

plt.xticks(rotation=25, ha="right")

# Add marks above bars

for bar, mark in zip(bars, marks):
    plt.text(
        bar.get_x() + bar.get_width() / 2,
        bar.get_height() + 2,
        f"{mark:.0f}%",
        ha="center"
    )

plt.tight_layout()

bar_chart_path = os.path.join(
    output_folder,
    "subject_performance.png"
)

plt.savefig(bar_chart_path, dpi=200)

plt.close()


# ============================================================
# 8. PIE CHART
# ============================================================

plt.figure(figsize=(8, 8))

plt.pie(
    marks,
    labels=subjects,
    autopct="%1.1f%%",
    startangle=90
)

plt.title("Marks Distribution")

plt.tight_layout()

pie_chart_path = os.path.join(
    output_folder,
    "marks_distribution.png"
)

plt.savefig(pie_chart_path, dpi=200)

plt.close()


# ============================================================
# 9. RADAR CHART
# ============================================================

angles = np.linspace(
    0,
    2 * np.pi,
    len(subjects),
    endpoint=False
)

# Close the radar chart

radar_marks = marks + [marks[0]]
radar_angles = list(angles) + [angles[0]]

fig = plt.figure(figsize=(8, 8))

ax = fig.add_subplot(111, polar=True)

ax.plot(
    radar_angles,
    radar_marks,
    linewidth=2
)

ax.fill(
    radar_angles,
    radar_marks,
    alpha=0.25
)

ax.set_xticks(angles)
ax.set_xticklabels(subjects)

ax.set_ylim(0, 100)

ax.set_title(
    "Student Performance Radar",
    pad=20
)

radar_chart_path = os.path.join(
    output_folder,
    "performance_radar.png"
)

plt.tight_layout()

plt.savefig(
    radar_chart_path,
    dpi=200,
    bbox_inches="tight"
)

plt.close()


# ============================================================
# 10. CREATE VISUAL DASHBOARD
# ============================================================

fig = plt.figure(figsize=(16, 10))

fig.suptitle(
    f"{student_name} - First Semester ECE Performance",
    fontsize=20,
    fontweight="bold"
)


# ---------- BAR CHART ----------

ax1 = plt.subplot(2, 2, 1)

bars = ax1.bar(subjects, marks)

ax1.set_title("Subject Performance")
ax1.set_ylabel("Percentage")
ax1.set_ylim(0, 100)

ax1.tick_params(axis="x", rotation=25)

for bar, mark in zip(bars, marks):

    ax1.text(
        bar.get_x() + bar.get_width() / 2,
        mark + 2,
        f"{mark:.0f}%",
        ha="center"
    )


# ---------- PIE CHART ----------

ax2 = plt.subplot(2, 2, 2)

ax2.pie(
    marks,
    labels=subjects,
    autopct="%1.1f%%",
    startangle=90
)

ax2.set_title("Marks Distribution")


# ---------- RADAR CHART ----------

ax3 = plt.subplot(2, 2, 3, polar=True)

ax3.plot(
    radar_angles,
    radar_marks,
    linewidth=2
)

ax3.fill(
    radar_angles,
    radar_marks,
    alpha=0.25
)

ax3.set_xticks(angles)

ax3.set_xticklabels(subjects)

ax3.set_ylim(0, 100)

ax3.set_title(
    "Performance Radar",
    pad=20
)


# ---------- SUMMARY ----------

ax4 = plt.subplot(2, 2, 4)

ax4.axis("off")

summary_text = (
    f"STUDENT: {student_name}\n\n"
    f"SEMESTER: {semester}\n\n"
    f"TOTAL: {total_marks:.0f}/{maximum_marks}\n\n"
    f"OVERALL: {overall_percentage:.2f}%\n\n"
    f"SUBJECTS: {number_of_subjects}\n\n"
    "Performance Summary\n"
    "-------------------\n"
)

# Find highest subject

highest_index = marks.index(max(marks))

summary_text += (
    f"Highest: {subjects[highest_index]}\n"
    f"Score: {marks[highest_index]:.0f}%\n\n"
)

# Count grades

grade_counts = pd.Series(grades).value_counts()

for grade, count in grade_counts.items():

    summary_text += (
        f"{grade}: {count} subject(s)\n"
    )


ax4.text(
    0.1,
    0.9,
    summary_text,
    fontsize=13,
    verticalalignment="top"
)

dashboard_path = os.path.join(
    output_folder,
    "semester_dashboard.png"
)

plt.tight_layout()

plt.savefig(
    dashboard_path,
    dpi=200,
    bbox_inches="tight"
)

plt.close()


# ============================================================
# 11. CREATE EXCEL FILE
# ============================================================

# Make student name safe for Windows filenames
safe_student_name = "".join(
    c for c in student_name if c.isalnum() or c in (" ", "_", "-")
).strip()

if not safe_student_name:
    safe_student_name = "Student"

safe_student_name = safe_student_name.replace(" ", "_")

excel_path = os.path.join(
    output_folder,
    f"{safe_student_name}_Semester_Result.xlsx"
)
workbook = Workbook()

sheet = workbook.active

sheet.title = "Semester Result"


# Title

sheet["A1"] = "FIRST SEMESTER ECE RESULT"

sheet["A1"].font = Font(
    bold=True,
    size=18
)

sheet.merge_cells("A1:E1")

sheet["A1"].alignment = Alignment(
    horizontal="center"
)


# Student information

sheet["A3"] = "Student Name"
sheet["B3"] = student_name

sheet["A4"] = "Semester"
sheet["B4"] = semester

sheet["A5"] = "Total Marks"
sheet["B5"] = f"{total_marks:.0f}/{maximum_marks}"

sheet["A6"] = "Overall Percentage"
sheet["B6"] = f"{overall_percentage:.2f}%"


# Table headers

headers = [
    "S.No.",
    "Subject",
    "Marks",
    "Percentage",
    "Grade"
]

for column, header in enumerate(headers, start=1):

    cell = sheet.cell(
        row=8,
        column=column
    )

    cell.value = header

    cell.font = Font(
        bold=True
    )

    cell.fill = PatternFill(
        "solid",
        fgColor="D9EAF7"
    )

    cell.alignment = Alignment(
        horizontal="center"
    )


# Insert data

for row_index, row in enumerate(
    df.itertuples(index=False),
    start=9
):

    sheet.cell(
        row=row_index,
        column=1,
        value=row[0]
    )

    sheet.cell(
        row=row_index,
        column=2,
        value=row[1]
    )

    sheet.cell(
        row=row_index,
        column=3,
        value=row[2]
    )

    sheet.cell(
        row=row_index,
        column=4,
        value=row[3] / 100
    )

    sheet.cell(
        row=row_index,
        column=4
    ).number_format = "0.00%"

    sheet.cell(
        row=row_index,
        column=5,
        value=row[4]
    )


# Column widths

sheet.column_dimensions["A"].width = 10
sheet.column_dimensions["B"].width = 40
sheet.column_dimensions["C"].width = 15
sheet.column_dimensions["D"].width = 18
sheet.column_dimensions["E"].width = 12


# Add Excel bar chart

chart = BarChart()

chart.title = "Subject Marks"

chart.y_axis.title = "Marks"

chart.x_axis.title = "Subjects"

data_reference = Reference(
    sheet,
    min_col=3,
    min_row=8,
    max_row=8 + number_of_subjects
)

category_reference = Reference(
    sheet,
    min_col=2,
    min_row=9,
    max_row=8 + number_of_subjects
)

chart.add_data(
    data_reference,
    titles_from_data=True
)

chart.set_categories(
    category_reference
)

chart.height = 8

chart.width = 15

sheet.add_chart(
    chart,
    "G3"
)


workbook.save(excel_path)


# ============================================================
# 12. CREATE PDF CERTIFICATE
# ============================================================

pdf_path = os.path.join(
    output_folder,
    f"{safe_student_name}_Semester_Certificate.pdf"
)

document = SimpleDocTemplate(
    pdf_path,
    pagesize=A4,
    rightMargin=40,
    leftMargin=40,
    topMargin=40,
    bottomMargin=40
)

styles = getSampleStyleSheet()

title_style = ParagraphStyle(
    "CertificateTitle",
    parent=styles["Title"],
    alignment=TA_CENTER,
    fontSize=24,
    spaceAfter=15
)

subtitle_style = ParagraphStyle(
    "Subtitle",
    parent=styles["Normal"],
    alignment=TA_CENTER,
    fontSize=14,
    spaceAfter=20
)

normal_center = ParagraphStyle(
    "Center",
    parent=styles["Normal"],
    alignment=TA_CENTER,
    fontSize=12
)

story = []


story.append(
    Paragraph(
        "FIRST SEMESTER PERFORMANCE CERTIFICATE",
        title_style
    )
)

story.append(
    Paragraph(
        "Electronics and Communication Engineering (ECE)",
        subtitle_style
    )
)

story.append(
    Paragraph(
        f"This certificate recognizes the academic performance "
        f"recorded for <b>{student_name}</b> during "
        f"<b>{semester}</b>.",
        normal_center
    )
)

story.append(Spacer(1, 25))


# Certificate result

result_text = (
    f"<b>Overall Percentage: {overall_percentage:.2f}%</b><br/>"
    f"Total Marks: {total_marks:.0f} / {maximum_marks}"
)

story.append(
    Paragraph(
        result_text,
        normal_center
    )
)

story.append(Spacer(1, 25))


# Subject table

table_data = [
    [
        "S.No.",
        "Subject",
        "Marks",
        "Percentage",
        "Grade"
    ]
]

for row in df.itertuples(index=False):

    table_data.append([
        row[0],
        row[1],
        f"{row[2]:.0f}/100",
        f"{row[3]:.0f}%",
        row[4]
    ])


table = Table(
    table_data,
    colWidths=[
        40,
        220,
        70,
        75,
        55
    ]
)

table.setStyle(
    TableStyle([
        (
            "BACKGROUND",
            (0, 0),
            (-1, 0),
            colors.HexColor("#1F4E78")
        ),
        (
            "TEXTCOLOR",
            (0, 0),
            (-1, 0),
            colors.white
        ),
        (
            "FONTNAME",
            (0, 0),
            (-1, 0),
            "Helvetica-Bold"
        ),
        (
            "ALIGN",
            (0, 0),
            (-1, -1),
            "CENTER"
        ),
        (
            "VALIGN",
            (0, 0),
            (-1, -1),
            "MIDDLE"
        ),
        (
            "GRID",
            (0, 0),
            (-1, -1),
            0.5,
            colors.grey
        ),
        (
            "BACKGROUND",
            (0, 1),
            (-1, -1),
            colors.whitesmoke
        ),
        (
            "ROWBACKGROUNDS",
            (0, 1),
            (-1, -1),
            [
                colors.white,
                colors.HexColor("#EAF2F8")
            ]
        )
    ])
)

story.append(table)

story.append(Spacer(1, 25))


# Dashboard image

story.append(
    Paragraph(
        "<b>Performance Visualization</b>",
        subtitle_style
    )
)

story.append(
    Image(
        dashboard_path,
        width=500,
        height=310
    )
)

story.append(Spacer(1, 20))

story.append(
    Paragraph(
        "Generated automatically using Python.",
        normal_center
    )
)

story.append(
    Spacer(1, 20)
)

story.append(
    Paragraph(
        "This is a student performance report generated from "
        "the entered marks and is not an official institutional certificate.",
        normal_center
    )
)


document.build(story)


# ============================================================
# 13. FINAL MESSAGE
# ============================================================

print("\n")
print("=" * 60)
print("                 GENERATION COMPLETE")
print("=" * 60)

print("\nFiles created:")

print(f"\n1. Excel:")
print(f"   {excel_path}")

print(f"\n2. PDF Certificate:")
print(f"   {pdf_path}")

print(f"\n3. Dashboard:")
print(f"   {dashboard_path}")

print(f"\n4. Bar Chart:")
print(f"   {bar_chart_path}")

print(f"\n5. Pie Chart:")
print(f"   {pie_chart_path}")

print(f"\n6. Radar Chart:")
print(f"   {radar_chart_path}")

print("\nEverything has been saved inside:")
print(f"   {output_folder}/")

print("\nDone! 🚀")