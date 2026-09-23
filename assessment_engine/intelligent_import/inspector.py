"""Intelligent Import Inspector - V2.4 Phase 1.

The inspector *reads* an Excel workbook and describes its structure:
workbook metadata, per-sheet geometry, candidate header rows, per-column
hints and data-quality observations.

Guarantees (Phase 1):

* purely structural and deterministic - same input, same profile,
* never writes to storage and never touches academic data or
  :class:`~assessment_engine.importer.ImportPlan`,
* no network calls and no AI dependencies,
* every structural inference carries a confidence value; nothing is ever
  claimed to have final academic meaning.

Usage::

    profile = inspect_workbook("marks.xlsx")
    for sheet in profile.sheets:
        print(sheet.name, sheet.header_block)
"""

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from assessment_engine.intelligent_import.models import (
    CellRange,
    ColumnProfile,
    ColumnRoleCandidate,
    HeaderBlock,
    HeaderCandidate,
    LayoutHints,
    SheetProfile,
    StructuralWarning,
    WorkbookProfile,
)


class InspectorError(Exception):
    """The workbook could not be read or inspected."""


# ------------------------------------------------------------
# Header normalisation
# ------------------------------------------------------------

_NON_ALNUM_RE = re.compile(r"[^0-9A-Za-z]+")


def normalize_header_text(value) -> str:
    """Clean a header cell for later matching.

    Original text is preserved separately; this function produces the
    form used for structural matching.  ``" Student Name "`` becomes
    ``"student name"`` and ``"Roll No."`` becomes ``"roll no"``.
    """
    cleaned = _clean_text(value)
    return _NON_ALNUM_RE.sub(" ", cleaned).strip().lower()


def _clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _raw_text(value) -> str:
    """The original cell text, preserved verbatim (blank cells -> "")."""
    if _is_blank(value):
        return ""
    return str(value)


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_numeric(value) -> bool:
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        try:
            float(value.strip())
            return True
        except ValueError:
            return False
    return False


# ------------------------------------------------------------
# Structural lexicons (candidates only - never final meaning)
# ------------------------------------------------------------

# Each entry maps a role kind to (label, ordered list of (exact key,
# confidence)).  The first exact key that matches wins; keys are spelled
# out explicitly so the detector stays deterministic and conservative.
_IDENTITY_LEXICON = {
    "student_id": (
        "Student ID",
        [
            ("student id", 0.90), ("studentid", 0.90), ("student no", 0.90),
            ("student number", 0.90), ("student code", 0.90), ("id", 0.55),
        ],
    ),
    "roll_no": (
        "Roll number",
        [
            ("roll no", 0.97), ("rollno", 0.97), ("roll number", 0.97),
            ("roll num", 0.97), ("roll", 0.85),
        ],
    ),
    "admission_no": (
        "Admission number",
        [
            ("admission no", 0.93), ("adm no", 0.93), ("admission number", 0.93),
            ("admission", 0.70),
        ],
    ),
    "registration_no": (
        "Registration number",
        [
            ("reg no", 0.93), ("regno", 0.93), ("reg number", 0.93),
            ("register no", 0.91), ("register number", 0.91),
            ("registration no", 0.91), ("registration number", 0.91),
            ("reg", 0.55),
        ],
    ),
    "student_name": (
        "Student name",
        [
            ("student name", 0.95), ("student s name", 0.95),
            ("name of student", 0.95), ("student full name", 0.95),
            ("name", 0.80), ("student", 0.80), ("candidate", 0.75),
        ],
    ),
    "email": (
        "Email",
        [
            ("email", 0.96), ("email id", 0.96), ("email address", 0.96),
            ("e mail", 0.96), ("emailid", 0.96), ("mail", 0.70),
        ],
    ),
}

_ORG_LEXICON = {
    "class": (
        "Class",
        [
            ("class", 0.95), ("standard", 0.90), ("std", 0.90),
            ("class name", 0.90), ("class division", 0.90),
        ],
    ),
    "section": (
        "Section",
        [
            ("section", 0.95), ("sec", 0.90), ("section name", 0.90),
        ],
    ),
    "department": (
        "Department",
        [
            ("department", 0.95), ("dept", 0.90), ("department name", 0.90),
        ],
    ),
    "course_program": (
        "Course/Program",
        [
            ("course", 0.92), ("course name", 0.90), ("program", 0.92),
            ("programme", 0.92), ("branch", 0.88), ("stream", 0.88),
        ],
    ),
    "semester": (
        "Semester",
        [
            ("semester", 0.95), ("sem", 0.85), ("term", 0.85),
            ("semester name", 0.90),
        ],
    ),
    "year": (
        "Year",
        [
            ("academic year", 0.95), ("year", 0.90), ("ay", 0.80), ("a y", 0.80),
        ],
    ),
    "batch": (
        "Batch",
        [
            ("batch", 0.95), ("batch name", 0.90), ("batch year", 0.90),
        ],
    ),
}

_ACADEMIC_LEXICON = {
    "subject": (
        "Subject",
        [
            ("subject", 0.97), ("subjects", 0.97), ("subject name", 0.95),
        ],
    ),
    "assessment_name": (
        "Assessment name",
        [
            ("assessment name", 0.97), ("assessment", 0.95), ("assessment type", 0.90),
            ("exam", 0.90), ("exam name", 0.90), ("examination", 0.90),
        ],
    ),
    "marks_obtained": (
        "Marks obtained",
        [
            ("marks obtained", 0.97), ("marks scored", 0.97), ("marks", 0.97),
            ("mark", 0.90), ("score", 0.92), ("score obtained", 0.92),
            ("obtained", 0.92), ("obtained marks", 0.92),
        ],
    ),
    "max_marks": (
        "Maximum marks",
        [
            ("maximum marks", 0.97), ("max marks", 0.95), ("full marks", 0.95),
            ("max", 0.85),
        ],
    ),
    "grade": (
        "Grade",
        [
            ("grade", 0.95), ("grade obtained", 0.95), ("grade point", 0.90),
            ("gpa", 0.90),
        ],
    ),
    "result": (
        "Result",
        [
            ("result", 0.95), ("result status", 0.95), ("status", 0.85),
            ("pass fail", 0.90),
        ],
    ),
}

_ASSESSMENT_TOKEN_WORDS = {
    "fa", "sa", "pt", "ut", "ia", "see", "mid", "midterm", "internal",
    "external", "final", "practical", "assignment", "test", "mock",
}
_ASSESSMENT_PHRASES = ("end sem", "end semester")

_EMPTY_SHEET_WARNING = StructuralWarning(
    "empty_sheet", "sheet has no non-blank cells"
)


# ------------------------------------------------------------
# Per-row statistics and scoring
# ------------------------------------------------------------


@dataclass
class _RowStats:
    non_empty: int
    numeric: int
    text: int
    max_text_length: int
    fill_ratio: float
    numeric_ratio: float


def _compute_stats(cells: Tuple, width: int) -> _RowStats:
    non_empty = numeric = text = 0
    max_length = 0
    for value in cells:
        if _is_blank(value):
            continue
        non_empty += 1
        if _is_numeric(value):
            numeric += 1
        else:
            text += 1
            max_length = max(max_length, len(str(value)))
    fill_ratio = non_empty / width if width else 0.0
    numeric_ratio = numeric / non_empty if non_empty else 0.0
    return _RowStats(
        non_empty, numeric, text, max_length, fill_ratio, numeric_ratio
    )


def _is_header_row(stats: _RowStats) -> bool:
    """A structurally plausible header row.

    More than one filled cell, fills at least half the sheet width and is
    not dominated by numbers.
    """
    return (
        stats.non_empty >= 2
        and stats.fill_ratio >= 0.5
        and stats.numeric_ratio < 0.5
    )


def _header_score(stats: _RowStats) -> float:
    return round(
        stats.non_empty + stats.fill_ratio * 10 + (1.0 - stats.numeric_ratio) * 5,
        3,
    )


def _header_confidence(stats: _RowStats) -> float:
    confidence = 0.30
    confidence += min(stats.fill_ratio, 1.0) * 0.45
    confidence += (1.0 - stats.numeric_ratio) * 0.30
    if stats.max_text_length >= 80:
        confidence -= 0.25
    if stats.non_empty < 2:
        confidence -= 0.40
    return round(max(0.0, min(0.99, confidence)), 3)


# ------------------------------------------------------------
# Header block detection
# ------------------------------------------------------------


def _detect_candidates(rows: List[Tuple[int, Tuple]], width: int) -> List[HeaderCandidate]:
    candidates = []
    for row_no, cells in rows:
        stats = _compute_stats(cells, width)
        if not stats.non_empty or not _is_header_row(stats):
            continue
        candidates.append(
            HeaderCandidate(
                row_number=row_no,
                values=[_raw_text(cell) for cell in cells],
                normalized=[normalize_header_text(cell) for cell in cells],
                non_empty_count=stats.non_empty,
                score=_header_score(stats),
                confidence=_header_confidence(stats),
            )
        )
    return candidates


def _group_runs(candidates: List[HeaderCandidate]) -> List[List[HeaderCandidate]]:
    runs = []
    for candidate in candidates:
        if runs and candidate.row_number == runs[-1][-1].row_number + 1:
            runs[-1].append(candidate)
        else:
            runs.append([candidate])
    return runs


def _select_header_block(
    rows: List[Tuple[int, Tuple]], width: int
) -> Tuple[Optional[HeaderBlock], List[HeaderCandidate], List[List[HeaderCandidate]]]:
    """Return (chosen block, all candidates, grouped runs).

    The chosen block is the first run of consecutive header-like rows
    (top-down scan), extended only while the following row stays a pure
    text header row.  A single candidate row is almost always the answer.
    """
    candidates = _detect_candidates(rows, width)
    if not candidates:
        return None, candidates, []

    runs = _group_runs(candidates)

    chosen = [candidates[0]]
    top_non_empty = candidates[0].non_empty_count
    for candidate in candidates[1:]:
        if candidate.row_number != chosen[-1].row_number + 1:
            break
        stats = _compute_stats(rows[candidate.row_number - 1][1], width)
        if stats.numeric:
            break
        # A genuine second header row is sparser than the row above it; a
        # full-width row below is more likely an all-text data row.
        if candidate.non_empty_count >= top_non_empty:
            break
        chosen.append(candidate)

    block_width = 0
    for candidate in chosen:
        for index, value in enumerate(candidate.values):
            if value and index + 1 > block_width:
                block_width = index + 1

    block = HeaderBlock(
        rows=[candidate.row_number for candidate in chosen],
        width=block_width,
        confidence=round(
            sum(candidate.confidence for candidate in chosen) / len(chosen), 3
        ),
    )
    return block, candidates, runs


# ------------------------------------------------------------
# Column analysis
# ------------------------------------------------------------


def _detect_roles(
    normalized: str,
) -> Tuple[List[ColumnRoleCandidate], List[str], List[str]]:
    roles = []
    if not normalized:
        return roles, [], []

    roles += _match_lexicon(normalized, _IDENTITY_LEXICON, "identity")
    roles += _match_lexicon(normalized, _ORG_LEXICON, "organizational")
    roles += _match_lexicon(normalized, _ACADEMIC_LEXICON, "academic")

    assessment_tokens, subject_tokens = _assessment_tokens(normalized)
    if assessment_tokens:
        confidence = min(0.95, 0.60 + 0.10 * len(set(assessment_tokens)))
        roles.append(
            ColumnRoleCandidate(
                family="assessment_like",
                kind="assessment_like",
                label="Assessment-like column",
                matched=", ".join(assessment_tokens),
                confidence=round(confidence, 3),
            )
        )
    return roles, assessment_tokens, subject_tokens


def _match_lexicon(normalized: str, lexicon: dict, family: str) -> List[ColumnRoleCandidate]:
    roles = []
    for kind, (label, keys) in lexicon.items():
        for key, confidence in keys:
            if normalized == key:
                roles.append(
                    ColumnRoleCandidate(
                        family=family,
                        kind=kind,
                        label=label,
                        matched=key,
                        confidence=confidence,
                    )
                )
                break
    return roles


def _assessment_tokens(normalized: str) -> Tuple[List[str], List[str]]:
    words = normalized.split()
    matched = []
    remaining = list(words)

    for phrase in _ASSESSMENT_PHRASES:
        if phrase in normalized:
            matched.append(phrase)

    phrase_words = set()
    for phrase in matched:
        for word in phrase.split():
            phrase_words.add(word)

    for word in words:
        if word in phrase_words:
            continue
        if word in _ASSESSMENT_TOKEN_WORDS and word not in matched:
            matched.append(word)

    subject_tokens = [
        word
        for word in words
        if word not in phrase_words and word not in _ASSESSMENT_TOKEN_WORDS
        and word.isalpha()
    ]
    return matched, subject_tokens


def _build_columns(header_values: List[str], width: int, data_rows) -> List[ColumnProfile]:
    columns = []
    for index in range(width):
        header_text = _raw_text(header_values[index]) if index < len(header_values) else ""
        normalized = normalize_header_text(header_text)

        non_empty = numeric = text = 0
        samples = []
        for _, cells in data_rows:
            if index >= len(cells):
                continue
            value = cells[index]
            if _is_blank(value):
                continue
            non_empty += 1
            if _is_numeric(value):
                numeric += 1
            else:
                text += 1
            if len(samples) < 5:
                samples.append(str(value))

        if non_empty == 0:
            value_type = "empty"
        elif numeric == non_empty:
            value_type = "numeric"
        elif text == non_empty:
            value_type = "text"
        else:
            value_type = "mixed"

        roles, assessment_tokens, subject_tokens = _detect_roles(normalized)
        subject_like = bool(
            header_text
            and not roles
            and non_empty > 0
            and value_type in ("numeric", "mixed")
        )

        columns.append(
            ColumnProfile(
                index=index,
                header_text=header_text,
                normalized=normalized,
                value_type=value_type,
                non_empty_count=non_empty,
                numeric_count=numeric,
                text_count=text,
                samples=samples,
                roles=roles,
                assessment_tokens=assessment_tokens,
                subject_tokens=subject_tokens,
                subject_like=subject_like,
            )
        )
    return columns


def _effective_header(rows_by_no: Dict[int, Tuple], block: HeaderBlock) -> List[str]:
    effective = [""] * block.width
    for row_no in block.rows:
        cells = rows_by_no.get(row_no, ())
        for index in range(min(block.width, len(cells))):
            value = _raw_text(cells[index])
            if value:
                effective[index] = value
    return effective


# ------------------------------------------------------------
# Layout hints and warnings
# ------------------------------------------------------------


def _build_layout_hints(columns: List[ColumnProfile]) -> LayoutHints:
    assessment_like = [col for col in columns if col.is_assessment_like]
    subject_like = [col for col in columns if col.subject_like]
    has_subject = any(col.has_role("subject") for col in columns)
    has_marks = any(col.has_role("marks_obtained") for col in columns)
    has_assessment_name = any(col.has_role("assessment_name") for col in columns)

    notes = []
    if len(assessment_like) > 1:
        notes.append(f"{len(assessment_like)} assessment-like columns detected")

    if has_subject and has_marks:
        orientation, confidence = "long", 0.95
    elif subject_like or len(assessment_like) > 1:
        orientation, confidence = "wide", 0.85
    elif len(columns) >= 4 or has_assessment_name or has_marks:
        orientation, confidence = "wide", 0.60
    else:
        orientation, confidence = "unknown", 0.20

    return LayoutHints(
        orientation=orientation,
        orientation_confidence=round(confidence, 3),
        multiple_assessment_columns=len(assessment_like) > 1,
        subject_like_columns=[column.index for column in subject_like],
        notes=notes,
    )


def _has_academic_columns(columns: List[ColumnProfile]) -> bool:
    return any(
        column.has_role("subject")
        or column.has_role("marks_obtained")
        or column.has_role("assessment_name")
        or column.has_role("grade")
        or column.has_role("max_marks")
        or column.subject_like
        or column.is_assessment_like
        for column in columns
    )


def _collect_warnings(
    columns: List[ColumnProfile],
    block: Optional[HeaderBlock],
    candidates: List[HeaderCandidate],
    runs: List[List[HeaderCandidate]],
    header_found: bool,
    data_row_count: int,
) -> List[StructuralWarning]:
    warnings = []

    if block and block.is_multi_row:
        warnings.append(
            StructuralWarning(
                "multi_row_header",
                f"rows {block.rows} form a multi-row header block; "
                "semantic merging is left to later phases",
                severity="info",
                row_number=block.first_row,
            )
        )

    if candidates and runs:
        first_run = runs[0]
        alternative_runs = [run for run in runs if run is not first_run]
        first_max = max(candidate.confidence for candidate in first_run)
        if alternative_runs:
            alt_max = max(
                max(candidate.confidence for candidate in run)
                for run in alternative_runs
            )
            warnings.append(
                StructuralWarning(
                    "multiple_possible_header_rows",
                    f"{len(alternative_runs)} other header candidate row(s) found "
                    f"(e.g. row {alternative_runs[0][0].row_number}); "
                    "the header choice may be wrong",
                    row_number=block.first_row if block else None,
                )
            )
            if abs(first_max - alt_max) <= 0.1:
                warnings.append(
                    StructuralWarning(
                        "ambiguous_header_region",
                        "two or more rows are equally plausible headers",
                        row_number=block.first_row if block else None,
                    )
                )
        elif first_max < 0.6:
            warnings.append(
                StructuralWarning(
                    "ambiguous_header_region",
                    "best header candidate has low confidence",
                    row_number=block.first_row if block else None,
                )
            )
    elif not header_found and block:
        warnings.append(
            StructuralWarning(
                "ambiguous_header_region",
                "no clear header row was found; the first non-blank row "
                "is assumed as a low-confidence header",
                row_number=block.first_row,
            )
        )

    seen = {}
    duplicates = []
    for column in columns:
        if not column.normalized:
            continue
        if column.normalized in seen:
            duplicates.append(column.header_text)
        else:
            seen[column.normalized] = True
    if duplicates:
        warnings.append(
            StructuralWarning(
                "duplicate_headers",
                "duplicate headers: " + ", ".join(duplicates),
            )
        )

    if columns and not any(column.is_identity for column in columns):
        warnings.append(
            StructuralWarning(
                "missing_identity_columns",
                "no candidate student identity column found "
                "(e.g. 'Roll No'/'Student Name'); exact meanings are not assumed",
            )
        )

    if columns and not _has_academic_columns(columns):
        warnings.append(
            StructuralWarning(
                "no_academic_columns",
                "no obvious academic (marks) column detected",
            )
        )

    if data_row_count < 2:
        warnings.append(
            StructuralWarning(
                "sparse_sheet",
                f"only {data_row_count} data row(s); sheet looks suspiciously sparse",
            )
        )

    return warnings


# ------------------------------------------------------------
# Public API
# ------------------------------------------------------------


def _load_rows(worksheet) -> List[Tuple[int, Tuple]]:
    rows = []
    for row_no, cells in enumerate(
        worksheet.iter_rows(values_only=True), start=1
    ):
        rows.append((row_no, tuple(cells)))
    return rows


def inspect_sheet(worksheet, sheet_name: Optional[str] = None) -> SheetProfile:
    """Build a :class:`SheetProfile` for one openpyxl worksheet.

    Purely structural; reads cell values only.
    """
    name = sheet_name or worksheet.title
    rows = _load_rows(worksheet)
    row_count = max((row_no for row_no, _ in rows), default=0)
    width = max((len(cells) for _, cells in rows), default=0)

    blank_rows = [row_no for row_no, cells in rows if not any(
        not _is_blank(cell) for cell in cells
    )]
    content_rows = [(row_no, cells) for row_no, cells in rows if any(
        not _is_blank(cell) for cell in cells
    )]

    if not content_rows:
        return SheetProfile(
            name=name,
            row_count=row_count,
            column_count=width,
            empty=True,
            blank_rows=blank_rows,
            warnings=[_EMPTY_SHEET_WARNING],
        )

    min_row = content_rows[0][0]
    max_row = content_rows[-1][0]
    min_col = min(
        index
        for _, cells in content_rows
        for index, cell in enumerate(cells)
        if not _is_blank(cell)
    )
    max_col = max(
        index
        for _, cells in content_rows
        for index, cell in enumerate(cells)
        if not _is_blank(cell)
    )
    non_empty_region = CellRange(min_row, min_col, max_row, max_col)

    block, candidates, runs = _select_header_block(rows, width)

    header_found = block is not None and block.rows
    if block is None or not block.rows:
        fallback_row, fallback_cells = content_rows[0]
        fallback_width = max(
            (
                index + 1
                for index, cell in enumerate(fallback_cells)
                if not _is_blank(cell)
            ),
            default=1,
        )
        block = HeaderBlock(
            rows=[fallback_row], width=fallback_width, confidence=0.40
        )

    rows_by_no = {row_no: cells for row_no, cells in rows}
    header_values = _effective_header(rows_by_no, block)

    data_content_rows = [
        (row_no, cells)
        for row_no, cells in content_rows
        if row_no > block.last_row
    ]

    data_region = None
    if data_content_rows:
        d_min_row = data_content_rows[0][0]
        d_max_row = data_content_rows[-1][0]
        d_min_col = min(
            index
            for _, cells in data_content_rows
            for index, cell in enumerate(cells)
            if not _is_blank(cell)
        )
        d_max_col = max(
            index
            for _, cells in data_content_rows
            for index, cell in enumerate(cells)
            if not _is_blank(cell)
        )
        data_region = CellRange(d_min_row, d_min_col, d_max_row, d_max_col)

    analysis_width = max(
        block.width, (data_region.max_col + 1) if data_region else 0
    )
    columns = _build_columns(header_values, analysis_width, data_content_rows)
    layout_hints = _build_layout_hints(columns)
    warnings = _collect_warnings(
        columns,
        block,
        candidates,
        runs,
        header_found=header_found,
        data_row_count=len(data_content_rows),
    )

    return SheetProfile(
        name=name,
        row_count=row_count,
        column_count=width,
        empty=False,
        non_empty_region=non_empty_region,
        blank_rows=blank_rows,
        header_candidates=candidates,
        header_block=block,
        data_region=data_region,
        columns=columns,
        layout_hints=layout_hints,
        warnings=warnings,
    )


def inspect_workbook(path) -> WorkbookProfile:
    """Inspect an Excel workbook and return its structural profile.

    Raises :class:`FileNotFoundError` when the path does not exist and
    :class:`InspectorError` when the file cannot be read as a workbook.
    """
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {workbook_path}")

    try:
        workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    except Exception as exc:
        raise InspectorError(f"unable to read workbook '{workbook_path}': {exc}") from exc

    try:
        sheets = [inspect_sheet(worksheet) for worksheet in workbook.worksheets]
    finally:
        workbook.close()

    return WorkbookProfile(
        source=str(workbook_path),
        filename=workbook_path.name,
        sheet_names=[sheet.name for sheet in sheets],
        sheets=sheets,
    )