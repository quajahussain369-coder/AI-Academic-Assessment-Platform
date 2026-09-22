"""V2.4 Intelligent Normalizer - Phase 3: canonical representation.

The normalizer consumes a Phase 1 :class:`WorkbookProfile` and a Phase 2
:class:`WorkbookDetection` and produces a canonical, deterministic
:class:`WorkbookNormalization`:

* student identity, organization, academic-period, subject, assessment,
  mark, percentage, grade, total and status fields are normalized per
  record using the Phase 2 semantic reading of each column,
* every normalized value keeps its raw source value and full provenance
  (workbook, sheet, row, column and header path),
* phase-2 ambiguities are preserved as :class:`Unresolved` records and
  never guessed.

Guarantees (Phase 3):

* deterministic and offline - same profile + detection, same normalization,
* consumes ``WorkbookProfile`` + ``WorkbookDetection``; it never modifies
  the importer, ``ImportPlan`` or any academic result,
* institution-agnostic - no hardcoded institution, subject, sheet name or
  filename,
* percentage fractions are only converted using semantic context,
* status tokens (AB/Absent/Exempt) are never turned into numerical zeros.

Usage::

    from assessment_engine.intelligent_import import (
        inspect_workbook, detect_workbook, normalize_workbook_profile,
    )

    profile = inspect_workbook("marks.xlsx")
    detection = detect_workbook(profile)
    normalization = normalize_workbook_profile(profile, detection=detection)
    for sheet in normalization.sheets:
        for record in sheet.records:
            print(record.subject, record.subject_mark)
"""

import re
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook

from assessment_engine.intelligent_import.detector.detector import detect_workbook
from assessment_engine.intelligent_import.detector.models import (
    ColumnSemantics,
    SemanticCandidate,
    SheetSemantics,
    WorkbookDetection,
)
from assessment_engine.intelligent_import.inspector import (
    inspect_workbook,
    normalize_header_text,
)
from assessment_engine.intelligent_import.models import SheetProfile
from assessment_engine.intelligent_import.normalizer.models import (
    CellLocation,
    NormalizedAssessment,
    NormalizedField,
    NormalizedRecord,
    NormalizedSheet,
    NormalizedValue,
    Unresolved,
    VALUE_TEXT,
    VALUE_UNRESOLVED,
    WorkbookNormalization,
)
from assessment_engine.intelligent_import.normalizer.value import (
    is_blank,
    normalize_grade,
    normalize_identifier,
    normalize_mark,
    normalize_measure,
    normalize_percentage,
    normalize_result,
    normalize_text,
    normalize_total,
)

# ------------------------------------------------------------
# Semantic-kind -> record-container routing
# ------------------------------------------------------------

_IDENTITY_KINDS = {
    "roll_no", "student_id", "admission_no", "registration_no",
    "student_name", "email", "serial_no",
}
_ORGANIZATION_KINDS = {"class", "section", "department", "course_program"}
_PERIOD_KINDS = {"semester", "year", "batch"}
_MEASURE_KINDS = {
    "total_marks", "marks_obtained", "percentage", "grade",
    "average", "rank", "attendance", "max_marks",
}

_TARGET_IDENTITY = "identity"
_TARGET_ORGANIZATION = "organization"
_TARGET_PERIOD = "period"
_TARGET_SUBJECT = "subject"
_TARGET_SUBJECT_MARK = "subject_mark"
_TARGET_ASSESSMENT = "assessment"
_TARGET_MEASURES = "measures"
_TARGET_STATUSES = "statuses"


def _target_for(kind: str) -> str:
    if kind in _IDENTITY_KINDS:
        return _TARGET_IDENTITY
    if kind in _ORGANIZATION_KINDS:
        return _TARGET_ORGANIZATION
    if kind in _PERIOD_KINDS:
        return _TARGET_PERIOD
    if kind == "subject":
        return _TARGET_SUBJECT
    if kind == "subject_mark":
        return _TARGET_SUBJECT_MARK
    if kind in ("assessment", "assessment_name", "subject_assessment"):
        return _TARGET_ASSESSMENT
    if kind == "result":
        return _TARGET_STATUSES
    return _TARGET_MEASURES


def _value_for_kind(kind: str, raw) -> NormalizedValue:
    """Normalize a cell by the Phase 2 semantic kind of its column."""
    if kind in (
        "subject", "student_name", "class", "section", "department",
        "course_program", "semester", "year", "batch", "email",
    ):
        return normalize_text(raw)
    if kind == "result":
        return normalize_result(raw)
    if kind in (
        "roll_no", "student_id", "admission_no", "registration_no",
        "serial_no",
    ):
        return normalize_identifier(raw)
    if kind in ("subject_mark", "marks_obtained", "assessment", "subject_assessment"):
        return normalize_mark(raw)
    if kind == "max_marks":
        return normalize_measure(raw, label="max marks")
    if kind == "percentage":
        return normalize_percentage(raw)
    if kind == "grade":
        return normalize_grade(raw)
    if kind == "total_marks":
        return normalize_total(raw)
    if kind in ("average", "rank", "attendance"):
        return normalize_measure(raw, label=kind)
    return normalize_measure(raw, label=kind or "value")


# ------------------------------------------------------------
# Assessment construction (header-derived names/subjects)
# ------------------------------------------------------------


def _recover_term(text: str, term: str) -> str:
    """Recover the originally-spelled form of ``term`` inside ``text``."""
    if not text or not term:
        return ""
    try:
        match = re.search(
            r"\b" + re.escape(term.strip()) + r"\b", str(text), re.IGNORECASE
        )
    except re.error:
        return ""
    return match.group(0) if match else ""


def _strip_term(text: str, term: str) -> str:
    """Remove the first occurrence of ``term`` from ``text``."""
    cleaned = str(text)
    display = _recover_term(cleaned, term)
    if display:
        try:
            cleaned = re.sub(
                r"\b" + re.escape(display) + r"\b",
                " ",
                cleaned,
                count=1,
                flags=re.IGNORECASE,
            )
        except re.error:
            pass
    return " ".join(cleaned.split())


def _assessment_display(header_path: List[str], header_text: str, term: str) -> str:
    if len(header_path) >= 2:
        return header_path[-1]
    return _recover_term(header_text, term) or (term or header_text)


def _subject_from_header(
    header_path: List[str],
    header_text: str,
    subject_norm: Optional[str],
    term: str,
) -> tuple:
    """Recover ``(display, normalized_subject)`` from the header hierarchy."""
    if len(header_path) >= 2:
        display = header_path[-2]
        norm = subject_norm or normalize_header_text(display) or None
        return display, norm
    display = _strip_term(header_text, term)
    if display:
        norm = subject_norm or normalize_header_text(display) or None
        return display, norm
    return "", subject_norm or None


def _build_assessment(
    best: SemanticCandidate,
    column: ColumnSemantics,
    raw,
    data_loc: CellLocation,
    header_row: int,
) -> NormalizedAssessment:
    if best.kind == "assessment_name":
        name = normalize_text(raw)
        if name.is_unresolved or name.kind == "blank":
            name = NormalizedValue(
                VALUE_TEXT, raw, name.normalized, location=data_loc,
                note="assessment name read from the data cell",
            )
        name.location = data_loc
        name.note = "assessment name read from the data cell"
        return NormalizedAssessment(name=name, mark=_mark_value(raw, data_loc))

    term = best.assessment or best.normalized or ""
    header_location = CellLocation(
        data_loc.sheet,
        header_row if header_row else data_loc.row,
        column.column,
        list(column.header_path),
        column.header_text,
    )
    name = NormalizedValue(
        VALUE_TEXT,
        _assessment_display(column.header_path, column.header_text, term),
        term,
        location=header_location,
        note="assessment name taken from the column header",
    )
    subject_norm = best.subject if best.kind == "subject_assessment" else None
    subject_display, subject_normalized = _subject_from_header(
        column.header_path, column.header_text, subject_norm, term
    )
    subject = None
    if subject_normalized:
        subject = NormalizedValue(
            VALUE_TEXT,
            subject_display,
            subject_normalized,
            location=header_location,
            note="subject name taken from the header hierarchy",
        )
    return NormalizedAssessment(name=name, mark=_mark_value(raw, data_loc), subject=subject)


def _mark_value(raw, data_loc: CellLocation) -> NormalizedValue:
    value = normalize_mark(raw)
    value.location = data_loc
    return value


# ------------------------------------------------------------
# Record construction
# ------------------------------------------------------------


def _attach(record: NormalizedRecord, target: str, field: NormalizedField) -> None:
    if target == _TARGET_IDENTITY:
        record.identity.append(field)
    elif target == _TARGET_ORGANIZATION:
        record.organization.append(field)
    elif target == _TARGET_PERIOD:
        record.period.append(field)
    elif target == _TARGET_SUBJECT:
        if record.subject is None:
            record.subject = field
    elif target == _TARGET_SUBJECT_MARK:
        record.subject_mark = field
    elif target == _TARGET_STATUSES:
        record.statuses.append(field)
    else:
        record.measures.append(field)


def _reconcile_subjects(record: NormalizedRecord) -> None:
    """Wire assessments to the record subject and vice versa."""
    for assessment in record.assessments:
        if assessment.subject is None and record.subject is not None:
            assessment.subject = record.subject.value
    if record.subject is None:
        for assessment in record.assessments:
            if assessment.subject is not None:
                record.subject = NormalizedField("subject", assessment.subject)
                break


def _build_record(
    sheet_name: str,
    columns: List[ColumnSemantics],
    cells: tuple,
    row_no: int,
    ambiguous_columns: Dict[int, List],
    header_row: int,
) -> NormalizedRecord:
    record = NormalizedRecord(sheet=sheet_name, row=row_no)
    for column in columns:
        raw = cells[column.column] if column.column < len(cells) else None
        location = CellLocation(
            sheet_name,
            row_no,
            column.column,
            list(column.header_path),
            column.header_text,
        )
        if column.column in ambiguous_columns:
            codes = ", ".join(a.code for a in ambiguous_columns[column.column])
            kind = column.best.kind if column.best is not None else "unresolved"
            record.unresolved.append(
                NormalizedField(
                    kind,
                    NormalizedValue(
                        VALUE_UNRESOLVED,
                        raw,
                        None,
                        location=location,
                        note=(
                            "column meaning is unresolved ("
                            f"{codes}); see the workbook unresolved records"
                        ),
                    ),
                )
            )
            continue

        best = column.best
        if best is None:
            record.unresolved.append(
                NormalizedField(
                    "unresolved",
                    NormalizedValue(
                        VALUE_UNRESOLVED,
                        raw,
                        None,
                        location=location,
                        note="no semantic interpretation for this column",
                    ),
                )
            )
            continue

        target = _target_for(best.kind)
        if target == _TARGET_ASSESSMENT:
            record.assessments.append(
                _build_assessment(best, column, raw, location, header_row)
            )
            continue

        value = _value_for_kind(best.kind, raw)
        value.location = location
        _attach(record, target, NormalizedField(best.kind, value))

    _reconcile_subjects(record)
    return record


# ------------------------------------------------------------
# Sheet and workbook normalization
# ------------------------------------------------------------


def normalize_sheet(
    sheet_profile: SheetProfile,
    sheet_semantics: SheetSemantics,
    cells_by_row: Dict[int, tuple],
) -> NormalizedSheet:
    """Normalize one sheet from its profile, semantics and cell values."""
    normalized = NormalizedSheet(name=sheet_profile.name)
    if sheet_profile.empty:
        normalized.empty = True
        normalized.warnings.append("empty_sheet: sheet has no non-blank cells")
        return normalized
    if sheet_semantics.empty:
        normalized.empty = True
        normalized.warnings.append("empty_sheet: no semantic content detected")
        return normalized

    block = sheet_profile.header_block
    header_rows = list(block.rows) if block is not None and block.rows else []
    normalized.header_rows = header_rows
    normalized.warnings.extend(
        f"{warning.code}: {warning.message}" for warning in sheet_profile.warnings
    )
    normalized.warnings.extend(sheet_semantics.notes)

    columns = sorted(sheet_semantics.columns, key=lambda cs: cs.column)
    header_row = header_rows[-1] if header_rows else 0

    data_rows = []
    for row_no in sorted(cells_by_row):
        if header_row and row_no <= header_row:
            continue
        if any(not is_blank(cell) for cell in cells_by_row[row_no]):
            data_rows.append(row_no)

    ambiguous_columns: Dict[int, List] = {}
    for column in columns:
        if not column.ambiguity:
            continue
        ambiguous_columns[column.column] = list(column.ambiguity)
        for ambiguity in column.ambiguity:
            normalized.unresolved.append(
                Unresolved(
                    code=ambiguity.code,
                    message=ambiguity.message,
                    sheet=sheet_profile.name,
                    column=column.column,
                    header_text=column.header_text,
                    header_path=list(column.header_path),
                    options=list(ambiguity.options),
                    confidence_gap=ambiguity.confidence_gap,
                    locations=[
                        CellLocation(
                            sheet_profile.name,
                            row_no,
                            column.column,
                            list(column.header_path),
                            column.header_text,
                        )
                        for row_no in data_rows
                    ],
                )
            )

    for row_no in data_rows:
        normalized.records.append(
            _build_record(
                sheet_profile.name,
                columns,
                cells_by_row[row_no],
                row_no,
                ambiguous_columns,
                header_row,
            )
        )
    return normalized


def _load_rows_by_sheet(source: str) -> Dict[str, Dict[int, tuple]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    result = {}
    try:
        for worksheet in workbook.worksheets:
            rows = {}
            for row_no, cells in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            ):
                rows[row_no] = tuple(cells)
            result[worksheet.title] = rows
    finally:
        workbook.close()
    return result


def _find_sheet(detection: WorkbookDetection, name: str) -> Optional[SheetSemantics]:
    for sheet in detection.sheets:
        if sheet.name == name:
            return sheet
    return None


def _unique(items: List[str]) -> List[str]:
    return list(dict.fromkeys(items))


def normalize_workbook_profile(
    profile, *, detection: Optional[WorkbookDetection] = None
) -> WorkbookNormalization:
    """Normalize an inspected workbook (a ``WorkbookProfile``).

    ``detection`` may be supplied to avoid re-detecting; otherwise a
    Phase 2 detection is produced automatically.  The source workbook is
    reopened read-only to read the raw cell values - the profile itself
    intentionally stores structure only.
    """
    if detection is None:
        detection = detect_workbook(profile)

    source = profile.source
    if not source or not Path(source).is_file():
        raise FileNotFoundError(f"workbook file not found: {source}")

    cells = _load_rows_by_sheet(source)
    normalization = WorkbookNormalization(
        source=profile.source, filename=profile.filename
    )
    for sheet_profile in profile.sheets:
        semantics = _find_sheet(detection, sheet_profile.name)
        if semantics is None:
            continue
        sheet_normalization = normalize_sheet(
            sheet_profile, semantics, cells.get(sheet_profile.name, {})
        )
        normalization.sheets.append(sheet_normalization)
        normalization.records.extend(sheet_normalization.records)
        normalization.unresolved.extend(sheet_normalization.unresolved)
        normalization.warnings.extend(sheet_normalization.warnings)
    normalization.warnings = _unique(normalization.warnings)
    return normalization


def normalize_workbook(path, *, vocabulary=None, config=None) -> WorkbookNormalization:
    """Full pipeline convenience: inspect + detect + normalize one file.

    ``vocabulary`` / ``config`` are forwarded to the Phase 2 detector.
    """
    profile = inspect_workbook(path)
    detection = detect_workbook(profile, vocabulary=vocabulary, config=config)
    return normalize_workbook_profile(profile, detection=detection)