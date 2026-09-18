"""Excel report generation for the Student Grade Analyzer.

Reproduces the styled .xlsx file V1 created: a merged title, the learner
summary block, the marks table with percentage formatting, and an
embedded bar chart.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, Reference

from grade_analyzer import config


def write_excel_report(result, output_dir=config.OUTPUT_DIR):
    """Build and save the Excel report for a SemesterResult.

    The file is named ``{safe_name}_Semester_Result.xlsx`` inside the
    output folder, exactly like V1.  Returns the path of the saved file.
    """
    safe_name = config.make_safe_filename(result.student_name)

    excel_path = output_dir / (safe_name + config.EXCEL_FILENAME_SUFFIX)

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Semester Result"

    # ---------- Title ----------
    sheet["A1"] = "FIRST SEMESTER ECE RESULT"

    sheet["A1"].font = Font(bold=True, size=18)

    sheet.merge_cells("A1:E1")

    sheet["A1"].alignment = Alignment(horizontal="center")

    # ---------- Student information ----------
    sheet["A3"] = "Student Name"
    sheet["B3"] = result.student_name

    sheet["A4"] = "Semester"
    sheet["B4"] = result.semester

    sheet["A5"] = "Total Marks"
    sheet["B5"] = f"{result.total_marks:.0f}/{result.maximum_marks}"

    sheet["A6"] = "Overall Percentage"
    sheet["B6"] = f"{result.overall_percentage:.2f}%"

    # ---------- Table headers ----------
    headers = [
        "S.No.",
        "Subject",
        "Marks",
        "Percentage",
        "Grade",
    ]

    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(row=8, column=column)

        cell.value = header

        cell.font = Font(bold=True)

        cell.fill = PatternFill("solid", fgColor=config.EXCEL_HEADER_FILL)

        cell.alignment = Alignment(horizontal="center")

    # ---------- Insert data ----------
    for data_row, subject in enumerate(result.subjects, start=9):

        subject_index = data_row - 9

        sheet.cell(
            row=data_row,
            column=1,
            value=data_row - 8,  # S.No. starts at 1
        )

        sheet.cell(
            row=data_row,
            column=2,
            value=subject.name,
        )

        sheet.cell(
            row=data_row,
            column=3,
            value=subject.mark,
        )

        percentage_cell = sheet.cell(
            row=data_row,
            column=4,
            value=subject.mark / 100,
        )

        percentage_cell.number_format = "0.00%"

        sheet.cell(
            row=data_row,
            column=5,
            value=result.grades[subject_index],
        )

    # ---------- Column widths ----------
    for column, width in config.EXCEL_COLUMN_WIDTHS.items():
        sheet.column_dimensions[column].width = width

    # ---------- Add Excel bar chart ----------
    chart = BarChart()

    chart.title = "Subject Marks"

    chart.y_axis.title = "Marks"

    chart.x_axis.title = "Subjects"

    data_reference = Reference(
        sheet,
        min_col=3,
        min_row=8,
        max_row=8 + result.number_of_subjects,
    )

    category_reference = Reference(
        sheet,
        min_col=2,
        min_row=9,
        max_row=8 + result.number_of_subjects,
    )

    chart.add_data(data_reference, titles_from_data=True)

    chart.set_categories(category_reference)

    chart.height = 8

    chart.width = 15

    sheet.add_chart(chart, "G3")

    workbook.save(excel_path)

    return excel_path