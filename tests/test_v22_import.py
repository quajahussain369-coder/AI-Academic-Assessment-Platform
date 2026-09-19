"""Tests for the V2.2 Excel bulk import (strict by default)."""

import copy
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from assessment_engine import importer, models
from assessment_engine.analyzer import (
    compute_all_results,
    make_context,
    provision_institution,
)
from assessment_engine.cli import cmd_import
from assessment_engine.config import load_config
from assessment_engine.importer import (
    CellError,
    ImporterError,
    parse_mark_cell,
)
from assessment_engine.storage import JsonStorage

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHOOL_CONFIG = REPO_ROOT / "configs" / "school.json"

SCHOOL_ASSESSMENTS = ["FA-1", "FA-2", "SA-1", "SA-2"]


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def make_school_storage(tmp_path):
    config = load_config(SCHOOL_CONFIG)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)
    return config, storage


def write_workbook(path, sheets):
    """sheets: dict {sheet_name: list_of_rows}."""
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


def maths_headers():
    return ["Roll No", "Student Name"] + SCHOOL_ASSESSMENTS


def mark_counts(storage, institution_id):
    return len(storage.load_all(models.Mark, institution_id))


def gather_marks(storage, institution_id):
    return {
        (mark.student_id, mark.assessment_id): (
            mark.obtained,
            mark.status,
        )
        for mark in storage.load_all(models.Mark, institution_id)
    }


# ------------------------------------------------------------
# parse_mark_cell
# ------------------------------------------------------------


def test_parse_mark_cell_passes_values():
    assert parse_mark_cell(20.5, 100) == ("entered", 20.5)
    assert parse_mark_cell("78", 100) == ("entered", 78.0)


def test_parse_mark_cell_absent_and_exempt():
    for token in ("A", "AB", "absent", "Abs"):
        assert parse_mark_cell(token, 100) == ("absent", 0.0)
    for token in ("E", "EX", "exempt"):
        assert parse_mark_cell(token, 100) == ("exempt", 0.0)


def test_parse_mark_cell_blank_is_skipped():
    assert parse_mark_cell(None, 100) is None
    assert parse_mark_cell("", 100) is None


def test_parse_mark_cell_validates_range():
    with pytest.raises(CellError):
        parse_mark_cell(-1, 100)
    with pytest.raises(CellError):
        parse_mark_cell(101, 100)
    with pytest.raises(CellError):
        parse_mark_cell("oops", 100)


# ------------------------------------------------------------
# Native import
# ------------------------------------------------------------

NATIVE_ROWS = [
    ["6A01", "Aarav Sharma", 20, 22, 78, 85],
    ["6A02", "Bhavna Patel", 25, 24, 95, "A"],
]


def test_native_happy_path(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + NATIVE_ROWS},
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False

    report = plan.apply(storage)
    assert report.marks_written == 8  # 4 for s1, 4 for s2 (one absent)
    assert not report.skipped

    marks = gather_marks(storage, config.institution.id)
    assert len(marks) == 8

    context = make_context(config, storage)
    s1 = next(s for s in context.data.students if s.id == "s1")
    result = next(
        c for c in compute_all_results(context) if c.student_id == "s1"
    )
    maths = next(c for c in result.course_results if c.course_id == "maths")
    assert maths.percentage == pytest.approx(82.7)


def test_blank_rows_and_empty_sheets_are_notes(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "maths_a": [maths_headers()] + NATIVE_ROWS,
            "EmptySheet": [["", "", ""]],
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False
    assert any("EmptySheet" in note for note in plan.notes)


def test_absent_exempt_blank_handling(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "maths_a": [
                maths_headers(),
                ["6A01", "Aarav Sharma", "E", "A", "", 85],
            ],
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False
    plan.apply(storage)

    marks = gather_marks(storage, config.institution.id)
    fa1 = next(a for a in storage.load_all(
        models.Assessment, config.institution.id
    ) if a.offering_id == "maths_a" and a.name == "FA-1")
    assert marks[("s1", fa1.id)] == (0.0, "exempt")
    fa2 = next(a for a in storage.load_all(
        models.Assessment, config.institution.id
    ) if a.offering_id == "maths_a" and a.name == "FA-2")
    assert marks[("s1", fa2.id)] == (0.0, "absent")


# ------------------------------------------------------------
# Strict validation
# ------------------------------------------------------------


def test_strict_default_commits_nothing_on_errors(tmp_path):
    config, storage = make_school_storage(tmp_path)
    invalid_rows = [
        ["6A01", "Aarav Sharma", 20, 22, 78, 85],
        ["6A02", "Bhavna Patel", 25, 24, 95, 150],  # over max
    ]
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + invalid_rows},
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True
    assert mark_counts(storage, config.institution.id) == 0

    with pytest.raises(ImporterError):
        plan.apply(storage)
    assert mark_counts(storage, config.institution.id) == 0


def test_cli_strict_import_returns_nonzero_and_writes_nothing(tmp_path):
    config, storage = make_school_storage(tmp_path)
    invalid_rows = [
        ["6A01", "Aarav Sharma", 20, 22, 78, 85],
        ["6A02", "Bhavna Patel", 200, 24, 95, "A"],
    ]
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + invalid_rows},
    )

    class Args:
        config_path = str(SCHOOL_CONFIG)
        data = str(storage.data_dir)
        file = str(workbook_path)
        template = None
        offering = None
        legacy = False
        skip_invalid = False
        yes = True

    args = Args()
    result = cmd_import(args)
    assert result == 1
    assert mark_counts(storage, config.institution.id) == 0


def test_skip_invalid_commits_valid_records(tmp_path):
    config, storage = make_school_storage(tmp_path)
    mixed_rows = [
        ["6A01", "Aarav Sharma", 20, 22, 78, 85],
        ["6A02", "Bhavna Patel", 25, 24, 95, 150],
        ["6A03", "Rohan Mehta", 10, 10, 10, 10],
    ]
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + mixed_rows},
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True

    report = plan.apply(storage, skip_invalid=True)
    # s1 4 + s2 3 valid (SA-2 over max rejected) + s3 4 = 11
    assert report.marks_written == 11
    assert len(report.skipped) == 1
    assert "150" in report.skipped[0].message
    assert mark_counts(storage, config.institution.id) == 11


def test_cli_skip_invalid_does_not_prompt_with_yes(tmp_path):
    config, storage = make_school_storage(tmp_path)
    mixed_rows = [
        ["6A01", "Aarav Sharma", 20, 22, 78, 85],
        ["6A02", "Bhavna Patel", -5, 24, 95, "A"],
    ]
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + mixed_rows},
    )

    class Args:
        config_path = str(SCHOOL_CONFIG)
        data = str(storage.data_dir)
        file = str(workbook_path)
        template = None
        offering = None
        legacy = False
        skip_invalid = True
        yes = True

    args = Args()
    assert cmd_import(args) == 0
    assert mark_counts(storage, config.institution.id) == 7


def test_prompt_decline_commits_nothing(tmp_path, monkeypatch):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + NATIVE_ROWS[:1]},
    )

    class Args:
        config_path = str(SCHOOL_CONFIG)
        data = str(storage.data_dir)
        file = str(workbook_path)
        template = None
        offering = None
        legacy = False
        skip_invalid = False
        yes = False

    prompts = {"n"}

    def fake_input(_prompt):
        return prompts.pop() if prompts else "n"

    monkeypatch.setattr("builtins.input", fake_input)
    args = Args()
    assert cmd_import(args) == 2
    assert mark_counts(storage, config.institution.id) == 0


def test_prompt_accept_commits(tmp_path, monkeypatch):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + NATIVE_ROWS[:1]},
    )

    class Args:
        config_path = str(SCHOOL_CONFIG)
        data = str(storage.data_dir)
        file = str(workbook_path)
        template = None
        offering = None
        legacy = False
        skip_invalid = False
        yes = False

    monkeypatch.setattr("builtins.input", lambda _prompt: "y")
    args = Args()
    assert cmd_import(args) == 0
    assert mark_counts(storage, config.institution.id) == 4


# ------------------------------------------------------------
# expect_course_prefix
# ------------------------------------------------------------

PREFIX_FALSE_CONFIG = {
    "institution": {
        "id": "sample_college",
        "name": "National College of Technology",
        "institution_type": "College",
        "org_levels": ["Department", "Program"],
        "default_grade_scale_id": "college_standard",
        "default_rule_set_id": "college_rules",
    },
    "grade_scales": [
        {
            "id": "college_standard",
            "name": "scale",
            "bands": [{"min_percentage": 90, "grade": "A+"}, {"min_percentage": 0, "grade": "F"}],
        }
    ],
    "rule_sets": [
        {
            "id": "college_rules",
            "institution_id": "sample_college",
            "grade_scale_id": "college_standard",
            "pass_percentage": 40,
            "per_course_pass_percentage": 40,
            "aggregation": "weighted_percentage",
            "decimals": 2,
        }
    ],
    "assessment_schemes": [
        {
            "id": "internal_external",
            "name": "internal",
            "assessments": [
                {"name": "Internal", "max_marks": 50, "weight": 0.40, "order": 1},
                {"name": "External", "max_marks": 100, "weight": 0.60, "order": 2},
            ],
        }
    ],
    "reporting": {},
    "excel": {"expect_course_prefix": False},
    "seed": {
        "academic_years": [{"id": "CY2025", "name": "AY2025"}],
        "terms": [{"id": "semester_1", "name": "S1", "order": 1}],
        "org_units": [
            {"id": "dept_cse", "level": "Department", "name": "CSE"},
            {"id": "prog_cse", "level": "Program", "name": "B.Tech", "parent_id": "dept_cse"},
        ],
        "courses": [{"id": "dsa", "name": "Data Structures", "code": "CS201"}],
        "offerings": [
            {
                "id": "dsa_prog",
                "year_id": "CY2025",
                "term_id": "semester_1",
                "org_unit_id": "prog_cse",
                "course_id": "dsa",
                "assessment_scheme_id": "internal_external",
                "max_marks": 100,
            }
        ],
        "students": [{"id": "c1", "name": "Meera Iyer", "roll_no": "CSE01"}],
        "enrollments": [
            {
                "id": "ce1",
                "student_id": "c1",
                "year_id": "CY2025",
                "term_id": "semester_1",
                "org_unit_id": "prog_cse",
                "course_ids": ["dsa_prog"],
            }
        ],
    },
}


def _load_prefix_false_config(tmp_path):
    path = tmp_path / "college.json"
    path.write_text(json.dumps(PREFIX_FALSE_CONFIG), encoding="utf-8")
    config = load_config(path)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)
    return config, storage


def test_single_course_with_offering_hint(tmp_path):
    config, storage = _load_prefix_false_config(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "Marks": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["CSE01", "Meera Iyer", 30, 60],
            ]
        },
    )
    plan = importer.parse_workbook(
        config, storage, workbook_path, offering_hint="dsa_prog"
    )
    assert plan.has_errors is False
    plan.apply(storage)
    assert mark_counts(storage, config.institution.id) == 2


def test_single_course_missing_offering_is_error(tmp_path):
    config, storage = _load_prefix_false_config(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "Marks": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["CSE01", "Meera Iyer", 30, 60],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True
    assert any("--offering" in e.message for e in plan.errors)


def test_template_build_and_reimport(tmp_path):
    config, storage = _load_prefix_false_config(tmp_path)
    template = importer.build_template(config, storage, offering_hint="dsa_prog")
    template_path = tmp_path / "template.xlsx"
    template.save(template_path)

    workbook = load_workbook(template_path)
    sheet = workbook["Marks"]
    sheet.append(["CSE02", "Karthik Rao", 20, 40])
    workbook.save(template_path)

    plan = importer.parse_workbook(config, storage, template_path, offering_hint="dsa_prog")
    assert plan.has_errors is False
    report = plan.apply(storage)
    assert report.students_created == 1
    assert report.marks_written == 2


# ------------------------------------------------------------
# Student creation and conflicts
# ------------------------------------------------------------

LEGACY_MINIMAL_CONFIG = {
    "institution": {
        "id": "legacy_school",
        "name": "Legacy School",
        "institution_type": "School",
        "org_levels": ["Class", "Section"],
        "default_grade_scale_id": "g",
        "default_rule_set_id": "r",
    },
    "grade_scales": [
        {
            "id": "g",
            "name": "scale",
            "bands": [{"min_percentage": 90, "grade": "A+"}, {"min_percentage": 0, "grade": "F"}],
        }
    ],
    "rule_sets": [
        {
            "id": "r",
            "institution_id": "legacy_school",
            "grade_scale_id": "g",
            "pass_percentage": 40,
            "per_course_pass_percentage": 35,
            "aggregation": "weighted_percentage",
            "decimals": 2,
        }
    ],
    "assessment_schemes": [
        {
            "id": "final",
            "name": "single final",
            "assessments": [{"name": "Final", "max_marks": 100, "weight": 1.0, "order": 1}],
        }
    ],
    "reporting": {},
    "seed": {
        "academic_years": [{"id": "Y", "name": "Year"}],
        "terms": [{"id": "T", "name": "Term", "order": 1}],
        "org_units": [
            {"id": "class_6", "level": "Class", "name": "6"},
            {"id": "sec_a", "level": "Section", "name": "A", "parent_id": "class_6"},
        ],
        "courses": [
            {"id": "maths", "name": "Mathematics", "code": "MATH"},
            {"id": "science", "name": "Science", "code": "SCI"},
        ],
        "offerings": [
            {
                "id": "maths_a",
                "year_id": "Y",
                "term_id": "T",
                "org_unit_id": "sec_a",
                "course_id": "maths",
                "assessment_scheme_id": "final",
                "max_marks": 100,
            },
            {
                "id": "science_a",
                "year_id": "Y",
                "term_id": "T",
                "org_unit_id": "sec_a",
                "course_id": "science",
                "assessment_scheme_id": "final",
                "max_marks": 100,
            },
        ],
        "students": [{"id": "ls1", "name": "Lily Rose", "roll_no": "6A01"}],
        "enrollments": [
            {
                "id": "le1",
                "student_id": "ls1",
                "year_id": "Y",
                "term_id": "T",
                "org_unit_id": "sec_a",
                "course_ids": ["maths_a", "science_a"],
            }
        ],
    },
}


def _load_legacy_config(tmp_path):
    path = tmp_path / "legacy.json"
    path.write_text(json.dumps(LEGACY_MINIMAL_CONFIG), encoding="utf-8")
    config = load_config(path)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)
    return config, storage


def _strip_students(seed_dict):
    seed = copy.deepcopy(seed_dict)
    seed["students"] = []
    seed["enrollments"] = []
    return seed


def test_creates_students_and_enrollments(tmp_path):
    data = copy.deepcopy(LEGACY_MINIMAL_CONFIG)
    data["seed"] = _strip_students(data["seed"])
    path = tmp_path / "no_students.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    config = load_config(path)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)

    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "maths_a": [
                ["Roll No", "Student Name", "Final"],
                ["7A01", "New Student", 72],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False
    report = plan.apply(storage)
    assert report.students_created == 1
    assert report.enrollments_created == 1
    assert report.marks_written == 1

    students = storage.load_all(models.Student, config.institution.id)
    assert any(s.roll_no == "7A01" for s in students)
    enrollments = storage.load_all(models.Enrollment, config.institution.id)
    assert len(enrollments) == 1
    assert enrollments[0].course_ids == ["maths_a"]


def test_conflicting_roll_and_name_is_error(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "maths_a": [
                maths_headers(),
                ["6A01", "Bhavna Patel", 20, 22, 78, 85],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True
    assert any("resolves to" in e.message for e in plan.errors)


# ------------------------------------------------------------
# Never invent academic configuration
# ------------------------------------------------------------


def test_unknown_offering_is_error_and_writes_nothing(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "no_such_offering": [maths_headers(), ["6A01", "Aarav Sharma", 1, 2, 3, 4]],
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True
    assert any("no course offering" in e.message for e in plan.errors)
    assert storage.load_all(models.CourseOffering, config.institution.id)
    assert mark_counts(storage, config.institution.id) == 0


def test_unknown_assessment_is_error(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {
            "maths_a": [
                ["Roll No", "Student Name", "Bogus"],
                ["6A01", "Aarav Sharma", 5],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is True
    assert any("no assessment named" in e.message for e in plan.errors)


# ------------------------------------------------------------
# Idempotency
# ------------------------------------------------------------


def test_reimport_is_idempotent(tmp_path):
    config, storage = make_school_storage(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers()] + NATIVE_ROWS[:1]},
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False
    first = plan.apply(storage)
    second = plan.apply(storage)

    assert first.marks_written == 4
    assert second.marks_written == 4
    assert mark_counts(storage, config.institution.id) == 4


# ------------------------------------------------------------
# Legacy
# ------------------------------------------------------------


def test_legacy_long_layout(tmp_path):
    config, storage = _load_legacy_config(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "legacy_long.xlsx",
        {
            "Results": [
                ["Student", "Roll No", "Class", "Subject", "Marks"],
                ["Lily Rose", "6A01", "A", "Mathematics", 80],
                ["Lily Rose", "6A01", "A", "Science", 70],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path, legacy=True)
    assert plan.has_errors is False
    report = plan.apply(storage)
    assert report.marks_written == 2

    context = make_context(config, storage)
    result = next(
        r for r in compute_all_results(context) if r.student_id == "ls1"
    )
    assert result.overall_percentage == pytest.approx(75)


def test_legacy_wide_layout(tmp_path):
    config, storage = _load_legacy_config(tmp_path)
    workbook_path = write_workbook(
        tmp_path / "legacy_wide.xlsx",
        {
            "Results": [
                ["Student", "Roll No", "Class", "Mathematics", "Science"],
                ["Lily Rose", "6A01", "A", 80, 70],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path, legacy=True)
    assert plan.has_errors is False
    plan.apply(storage)
    assert mark_counts(storage, config.institution.id) == 2


def test_legacy_ambiguous_subject_is_error(tmp_path):
    data = copy.deepcopy(LEGACY_MINIMAL_CONFIG)
    data["seed"]["offerings"].append(copy.deepcopy(data["seed"]["offerings"][0]))
    data["seed"]["offerings"][2]["id"] = "maths_b"
    data["seed"]["offerings"][2]["org_unit_id"] = "class_6"
    path = tmp_path / "ambiguous.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    config = load_config(path)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)

    workbook_path = write_workbook(
        tmp_path / "legacy_long.xlsx",
        {
            "Results": [
                ["Student", "Roll No", "Class", "Subject", "Marks"],
                ["Lily Rose", "6A01", "", "Mathematics", 80],
            ]
        },
    )
    plan = importer.parse_workbook(config, storage, workbook_path, legacy=True)
    assert plan.has_errors is True
    assert any("ambiguous" in e.message for e in plan.errors)