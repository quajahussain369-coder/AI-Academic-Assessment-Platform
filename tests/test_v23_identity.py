"""V2.3 Step 2: identity & authorization foundation.

Verifies the platform-scoped ``User`` model, the institution-scoped
``Membership`` model, the pure authorization helpers in
``assessment_engine.auth``, platform-wide JSON user storage and the CLI's
optional ``--user`` enforcement.  No V2.1/V2.2/V2.3-Step-1 behaviour
changes.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook

from assessment_engine import auth, importer, models
from assessment_engine.analyzer import provision_institution
from assessment_engine.cli import cmd_import, cmd_init, cmd_list, cmd_user
from assessment_engine.config import load_config
from assessment_engine.storage import JsonStorage, Storage

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHOOL_CONFIG = REPO_ROOT / "configs" / "school.json"
COLLEGE_CONFIG = REPO_ROOT / "configs" / "college.json"

SCHOOL_ASSESSMENTS = ["FA-1", "FA-2", "SA-1", "SA-2"]

ADMIN1 = "admin1"


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def make_two_tenant_storage(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    provision_institution(load_config(SCHOOL_CONFIG), storage)
    provision_institution(load_config(COLLEGE_CONFIG), storage)
    return storage


def school_memberships():
    return [models.Membership("mem_admin1_school", ADMIN1, "sample_school", auth.ROLE_ADMIN)]


def add_admin_users(storage):
    storage.save_user(models.User(id=ADMIN1, name="Admin One", email="admin@example.com"))


def maths_headers():
    return ["Roll No", "Student Name"] + SCHOOL_ASSESSMENTS


def college_headers():
    return ["Roll No", "Student Name", "Internal", "External"]


def write_workbook(path, sheets):
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


class ImportArgs:
    def __init__(self, data, file, user=None, legacy=False):
        self.config_path = str(SCHOOL_CONFIG)
        self.data = data
        self.file = str(file)
        self.template = None
        self.offering = None
        self.legacy = legacy
        self.skip_invalid = False
        self.yes = True
        self.user = user


class AddUserArgs:
    action = "add"

    def __init__(self, data, user_id, name=None, email="", institution=None, role=None):
        self.data = data
        self.user_id = user_id
        self.name = name
        self.email = email
        self.institution = institution
        self.role = role


class ListUsersArgs:
    action = "list"

    def __init__(self, data, institution=None):
        self.data = data
        self.institution = institution


class InitArgs:
    def __init__(self, config_path, data, user=None):
        self.config_path = str(config_path)
        self.data = data
        self.user = user


# ------------------------------------------------------------
# User storage (platform-wide)
# ------------------------------------------------------------


def test_user_save_load_round_trip(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    storage.save_user(models.User(id="u1", name="Ada", email="ada@example.com"))
    loaded = storage.load_user("u1")
    assert loaded == models.User(id="u1", name="Ada", email="ada@example.com")


def test_load_missing_user_returns_none(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    assert storage.load_user("nobody") is None


def test_load_all_users_sorted(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    storage.save_user(models.User(id="zeta", name="Z"))
    storage.save_user(models.User(id="alpha", name="A"))
    assert [user.id for user in storage.load_all_users()] == ["alpha", "zeta"]


def test_users_are_platform_wide_and_shared_across_institutions(tmp_path):
    init_args = InitArgs(SCHOOL_CONFIG, str(tmp_path / "data"))
    init_args_2 = InitArgs(COLLEGE_CONFIG, str(tmp_path / "data"))
    assert cmd_init(init_args) == 0
    assert cmd_init(init_args_2) == 0

    storage = make_two_tenant_storage(tmp_path)
    storage.save_user(models.User(id=ADMIN1, name="Admin One", email="admin@example.com"))

    fresh = JsonStorage(tmp_path / "data")
    assert fresh.load_user(ADMIN1).name == "Admin One"
    assert (tmp_path / "data" / "users.json").is_file()

    for institution_id in ("sample_school", "sample_college"):
        database = storage._read(institution_id)
        assert "user" not in database


def test_users_json_does_not_interfere_with_list_institutions(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    storage.save_user(models.User(id="u1", name="U"))
    assert storage.list_institutions() == ["sample_college", "sample_school"]


def test_custom_storage_subclass_still_instantiable_without_user_methods():
    class LegacyStorage(Storage):
        def save(self, record):
            return None

        def load(self, model, obj_id, institution_id):
            return None

        def load_all(self, model, institution_id):
            return []

        def clear(self, institution_id):
            return None

    storage = LegacyStorage()
    assert isinstance(storage, Storage)
    with pytest.raises(NotImplementedError):
        storage.save_user(models.User(id="u", name="U"))
    with pytest.raises(NotImplementedError):
        storage.load_user("u")
    with pytest.raises(NotImplementedError):
        storage.load_all_users()


# ------------------------------------------------------------
# Membership storage (institution-scoped)
# ------------------------------------------------------------


def test_membership_save_load_round_trip(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    storage.save(models.Institution(id="inst", name="Inst"))
    storage.save(
        models.Membership("mem1", "u1", "inst", auth.ROLE_FACULTY)
    )
    loaded = storage.load(models.Membership, "mem1", "inst")
    assert loaded == models.Membership("mem1", "u1", "inst", auth.ROLE_FACULTY)


def test_membership_role_round_trip_through_cli(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    assert cmd_user(
        AddUserArgs(
            str(storage.data_dir), ADMIN1,
            institution="sample_school", role=auth.ROLE_STAFF,
        )
    ) == 0

    memberships = storage.load_all(models.Membership, "sample_school")
    assert len(memberships) == 1
    assert memberships[0].user_id == ADMIN1
    assert memberships[0].role == auth.ROLE_STAFF


def test_memberships_are_stored_per_institution(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    storage.save(
        models.Membership("mem1", ADMIN1, "sample_school", auth.ROLE_ADMIN)
    )
    assert len(storage.load_all(models.Membership, "sample_school")) == 1
    assert storage.load_all(models.Membership, "sample_college") == []


def test_duplicate_membership_upsert_updates_role_not_duplicates(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    for role in (auth.ROLE_STAFF, auth.ROLE_ADMIN, auth.ROLE_FACULTY):
        assert cmd_user(
            AddUserArgs(
                str(storage.data_dir), ADMIN1,
                institution="sample_school", role=role,
            )
        ) == 0
    memberships = storage.load_all(models.Membership, "sample_school")
    assert len(memberships) == 1
    assert memberships[0].id == f"mem_{ADMIN1}_sample_school"
    assert memberships[0].role == auth.ROLE_FACULTY


def test_adding_membership_preserves_existing_user_profile(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    assert cmd_user(
        AddUserArgs(
            str(storage.data_dir), ADMIN1,
            name="Admin One", email="admin@example.com",
        )
    ) == 0
    assert cmd_user(
        AddUserArgs(
            str(storage.data_dir), ADMIN1,
            institution="sample_school", role=auth.ROLE_ADMIN,
        )
    ) == 0
    user = storage.load_user(ADMIN1)
    assert user.name == "Admin One"
    assert user.email == "admin@example.com"


# ------------------------------------------------------------
# Authorization helpers (pure)
# ------------------------------------------------------------


def test_role_permission_mapping():
    assert auth.has_permission(auth.ROLE_ADMIN, auth.PERM_MANAGE_INSTITUTION)
    assert auth.has_permission(auth.ROLE_ADMIN, auth.PERM_MANAGE_USERS)
    assert auth.has_permission(auth.ROLE_FACULTY, auth.PERM_RECORD_MARKS)
    assert auth.has_permission(auth.ROLE_FACULTY, auth.PERM_VIEW_REPORTS)
    assert not auth.has_permission(auth.ROLE_FACULTY, auth.PERM_MANAGE_INSTITUTION)
    assert auth.has_permission(auth.ROLE_STAFF, auth.PERM_IMPORT)
    assert not auth.has_permission(auth.ROLE_STAFF, auth.PERM_RECORD_MARKS)
    assert auth.has_permission(auth.ROLE_STUDENT, auth.PERM_VIEW_OWN_RESULT)
    assert not auth.has_permission(auth.ROLE_STUDENT, auth.PERM_VIEW_REPORTS)
    assert not auth.has_permission("superuser", auth.PERM_MANAGE_INSTITUTION)


def test_is_valid_role():
    for role in auth.KNOWN_ROLES:
        assert auth.is_valid_role(role)
    assert not auth.is_valid_role("root")
    assert not auth.is_valid_role("")


def test_role_in_and_has_role():
    memberships = school_memberships()
    assert auth.role_in(memberships, ADMIN1, "sample_school") == auth.ROLE_ADMIN
    assert auth.role_in(memberships, ADMIN1, "sample_college") is None
    assert auth.role_in(memberships, "other", "sample_school") is None
    assert auth.has_role(memberships[0], auth.ROLE_ADMIN)
    assert not auth.has_role(memberships[0], auth.ROLE_STUDENT)


def test_authorize_true_for_valid_membership_and_permission():
    assert auth.authorize(
        school_memberships(), ADMIN1, "sample_school", auth.PERM_IMPORT
    )
    assert auth.authorize(
        school_memberships(), ADMIN1, "sample_school", auth.PERM_MANAGE_INSTITUTION
    )


def test_authorize_false_for_non_member():
    assert not auth.authorize(
        school_memberships(), ADMIN1, "sample_college", auth.PERM_IMPORT
    )
    assert not auth.authorize(
        school_memberships(), "stranger", "sample_school", auth.PERM_IMPORT
    )


def test_authorize_false_when_role_lacks_permission():
    student = [models.Membership("m", "stu", "sample_school", auth.ROLE_STUDENT)]
    assert auth.authorize(student, "stu", "sample_school", auth.PERM_VIEW_OWN_RESULT)
    assert not auth.authorize(student, "stu", "sample_school", auth.PERM_IMPORT)


def test_require_authorized_raises_authorization_error():
    with pytest.raises(auth.AuthorizationError):
        auth.require_authorized(
            school_memberships(), ADMIN1, "sample_college", auth.PERM_IMPORT
        )


def test_require_authorized_passes_for_allowed():
    auth.require_authorized(
        school_memberships(), ADMIN1, "sample_school", auth.PERM_IMPORT
    )


# ------------------------------------------------------------
# CLI: user management
# ------------------------------------------------------------


def test_cli_user_add_creates_platform_user(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    assert cmd_user(
        AddUserArgs(str(storage.data_dir), "ada", name="Ada L", email="ada@example.com")
    ) == 0
    user = storage.load_user("ada")
    assert user.name == "Ada L"
    assert user.email == "ada@example.com"
    assert user.status == "active"


def test_cli_user_list_lists_platform_users(tmp_path, capsys):
    storage = JsonStorage(tmp_path / "data")
    cmd_user(AddUserArgs(str(storage.data_dir), "ada", name="Ada", email="ada@e", role=None))
    cmd_user(AddUserArgs(str(storage.data_dir), "ben", name="Ben", email="ben@e", role=None))
    assert cmd_user(ListUsersArgs(str(storage.data_dir))) == 0
    captured = capsys.readouterr().out
    assert "ada" in captured
    assert "ben" in captured


def test_cli_user_add_requires_role_with_institution(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    assert (
        cmd_user(
            AddUserArgs(str(storage.data_dir), "u1", institution="sample_school")
        )
        == 2
    )


def test_cli_user_add_rejects_invalid_role(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    assert (
        cmd_user(
            AddUserArgs(
                str(storage.data_dir), "u1",
                institution="sample_school", role="superuser",
            )
        )
        == 2
    )


def test_cli_user_add_rejects_unknown_institution(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    assert (
        cmd_user(
            AddUserArgs(
                str(storage.data_dir), "u1",
                institution="nowhere", role=auth.ROLE_ADMIN,
            )
        )
        == 3
    )


def test_cli_user_list_memberships_by_institution(tmp_path, capsys):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    cmd_user(
        AddUserArgs(
            str(storage.data_dir), ADMIN1,
            institution="sample_school", role=auth.ROLE_ADMIN,
        )
    )
    assert cmd_user(ListUsersArgs(str(storage.data_dir), institution="sample_school")) == 0
    captured = capsys.readouterr().out
    assert ADMIN1 in captured
    assert auth.ROLE_ADMIN in captured
    assert "Members of 'sample_school'" in captured


# ------------------------------------------------------------
# CLI: --user enforcement
# ------------------------------------------------------------


def make_school_import_args(storage, tmp_path, user):
    workbook = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers(), ["6A01", "Aarav Sharma", 20, 22, 78, 85]]},
    )
    return ImportArgs(str(storage.data_dir), workbook, user=user)


def test_cli_authorized_operation_allowed(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    storage.save(
        models.Membership("mem1", ADMIN1, "sample_school", auth.ROLE_ADMIN)
    )
    args = make_school_import_args(storage, tmp_path, ADMIN1)
    assert cmd_import(args) == 0
    assert len(storage.load_all(models.Mark, "sample_school")) == 4


def test_cli_denied_operation_rejected_cross_institution(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    storage.save(
        models.Membership("mem1", ADMIN1, "sample_school", auth.ROLE_ADMIN)
    )

    before = len(storage.load_all(models.Mark, "sample_college"))
    workbook = write_workbook(
        tmp_path / "college.xlsx",
        {"dsa_prog": [college_headers(), ["CSE01", "Meera Iyer", 40, 80]]},
    )
    args = ImportArgs(str(storage.data_dir), workbook, user=ADMIN1)
    args.config_path = str(COLLEGE_CONFIG)
    assert cmd_import(args) == 3
    assert len(storage.load_all(models.Mark, "sample_college")) == before


def test_cli_denied_unknown_user(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    args = make_school_import_args(storage, tmp_path, "ghost")
    assert cmd_import(args) == 3


def test_cli_denied_inactive_user(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    storage.save_user(models.User(id="off", name="Off", status="inactive"))
    storage.save(
        models.Membership("mem1", "off", "sample_school", auth.ROLE_ADMIN)
    )
    args = make_school_import_args(storage, tmp_path, "off")
    assert cmd_import(args) == 3


def test_cli_denied_when_role_lacks_permission(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    storage.save(
        models.Membership("mem1", ADMIN1, "sample_school", auth.ROLE_STUDENT)
    )
    args = make_school_import_args(storage, tmp_path, ADMIN1)
    assert cmd_import(args) == 3


def test_no_user_preserves_existing_behavior(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    args = make_school_import_args(storage, tmp_path, user=None)
    assert cmd_import(args) == 0
    assert len(storage.load_all(models.Mark, "sample_school")) == 4


def test_cli_init_bootstrap_with_known_active_user(tmp_path):
    data_dir = str(tmp_path / "data")
    storage = JsonStorage(data_dir)
    add_admin_users(storage)
    assert cmd_init(InitArgs(SCHOOL_CONFIG, data_dir, user=ADMIN1)) == 0
    assert storage.list_institutions() == ["sample_school"]


def test_cli_init_denied_unknown_user(tmp_path):
    storage = JsonStorage(str(tmp_path / "data"))
    assert cmd_init(InitArgs(SCHOOL_CONFIG, str(tmp_path / "data"), user="ghost")) == 3
    assert storage.list_institutions() == []


def test_cli_init_denied_when_memberships_exist_and_user_not_admin(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    storage.save(
        models.Membership("mem1", ADMIN1, "sample_school", auth.ROLE_STUDENT)
    )
    assert cmd_init(InitArgs(SCHOOL_CONFIG, str(storage.data_dir), user=ADMIN1)) == 3


# ------------------------------------------------------------
# Documented provisioning behaviour
# ------------------------------------------------------------


def test_reinit_removes_memberships_documented_behaviour(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    add_admin_users(storage)
    cmd_user(
        AddUserArgs(
            str(storage.data_dir), ADMIN1,
            institution="sample_school", role=auth.ROLE_ADMIN,
        )
    )
    assert len(storage.load_all(models.Membership, "sample_school")) == 1

    assert cmd_init(InitArgs(SCHOOL_CONFIG, str(storage.data_dir))) == 0
    fresh = JsonStorage(str(storage.data_dir))
    assert len(fresh.load_all(models.Membership, "sample_school")) == 0
    assert fresh.load_user(ADMIN1) is not None