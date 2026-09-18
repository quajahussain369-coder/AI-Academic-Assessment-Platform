"""End-to-end tests for the school-style and college-style example configs.

Each test provisions the configuration, records marks, computes results
and (for the school) generates real Excel/PDF report files into a
temporary folder.
"""

from pathlib import Path

import openpyxl
import pytest

from assessment_engine import models
from assessment_engine.analyzer import (
    compute_all_results,
    make_context,
    provision_institution,
    record_mark,
)
from assessment_engine.config import EXCEL_SUFFIX, PDF_SUFFIX, load_config
from assessment_engine.reports import generate_reports
from assessment_engine.storage import JsonStorage

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHOOL_CONFIG = REPO_ROOT / "configs" / "school.json"
COLLEGE_CONFIG = REPO_ROOT / "configs" / "college.json"


def provision(config_path, tmp_path):
    config = load_config(config_path)
    storage = JsonStorage(tmp_path / "data")
    provision_institution(config, storage)
    return config, storage


def assessments_by_offering(storage, institution_id):
    by_name = {}
    for assessment in storage.load_all(models.Assessment, institution_id):
        by_name.setdefault(assessment.offering_id, {})[assessment.name] = assessment
    return by_name


def record_scores(storage, institution_id, assessments_by_offering_name, scores):
    """scores maps student_id -> offering_id -> {assessment_name: value}."""
    for student_id, offerings in scores.items():
        for offering_id, assessment_scores in offerings.items():
            for name, value in assessment_scores.items():
                assessment = assessments_by_offering_name[offering_id][name]
                obtained, status = value if isinstance(value, tuple) else (value, "entered")
                record_mark(
                    storage, institution_id, student_id, assessment.id,
                    obtained, status,
                )


SCHOOL_SCORES = {
    "s1": {
        "maths_a": {"FA-1": 20, "FA-2": 22, "SA-1": 78, "SA-2": 85},
        "science_a": {"FA-1": 18, "FA-2": 20, "SA-1": 62, "SA-2": 70},
        "english_a": {"FA-1": 21, "FA-2": 19, "SA-1": 66, "SA-2": 58},
    },
    "s2": {
        "maths_a": {"FA-1": 25, "FA-2": 24, "SA-1": 95, "SA-2": (0, "absent")},
        "science_a": {"FA-1": 22, "FA-2": 24, "SA-1": 88, "SA-2": 92},
        "english_a": {"FA-1": 20, "FA-2": 23, "SA-1": 78, "SA-2": 80},
    },
    "s3": {
        "maths_a": {"FA-1": 10, "FA-2": 8, "SA-1": 30, "SA-2": 25},
        "science_a": {"FA-1": 12, "FA-2": 10, "SA-1": 28, "SA-2": 20},
        "english_a": {"FA-1": 14, "FA-2": 12, "SA-1": 40, "SA-2": 25},
    },
}


def make_school_results(config, storage):
    by_offering = assessments_by_offering(storage, config.institution.id)
    record_scores(storage, config.institution.id, by_offering, SCHOOL_SCORES)
    context = make_context(config, storage)
    return {result.student_id: result for result in compute_all_results(context)}


# ------------------------------------------------------------
# School-style configuration
# ------------------------------------------------------------


def test_school_provisions_scheme_and_structure(tmp_path):
    config, storage = provision(SCHOOL_CONFIG, tmp_path)
    assessments = storage.load_all(models.Assessment, config.institution.id)
    assert len(assessments) == 12

    org_units = storage.load_all(models.OrgUnit, config.institution.id)
    assert {unit.level for unit in org_units} == {"Class", "Section"}


def test_school_arav_sharma_weighted_result(tmp_path):
    config, storage = provision(SCHOOL_CONFIG, tmp_path)
    results = make_school_results(config, storage)

    result = results["s1"]
    assert result.student_name == "Aarav Sharma"
    assert result.org_path == "Class 6 > Section A"
    assert result.number_of_courses == 3

    maths = next(c for c in result.course_results if c.course_id == "maths")
    assert maths.percentage == pytest.approx(82.7)
    assert maths.grade == "A"

    science = next(c for c in result.course_results if c.course_id == "science")
    assert science.percentage == pytest.approx(68.8)
    assert science.grade == "B"

    english = next(c for c in result.course_results if c.course_id == "english")
    assert english.percentage == pytest.approx(64.8)
    assert english.grade == "B"

    assert result.total_obtained == pytest.approx(216.3)
    assert result.total_maximum == pytest.approx(300)
    assert result.overall_percentage == pytest.approx(72.1)
    assert result.grade == "B+"
    assert result.passed is True


def test_school_absent_in_sa2_lowers_result(tmp_path):
    config, storage = provision(SCHOOL_CONFIG, tmp_path)
    results = make_school_results(config, storage)

    result = results["s2"]
    maths = next(c for c in result.course_results if c.course_id == "maths")
    assert maths.percentage == pytest.approx(48.1)
    assert maths.grade == "D"

    # The large weighted SA-2 contributed zero, so the overall mark drops.
    assert result.overall_percentage == pytest.approx(73.17)
    assert result.passed is True


def test_school_low_scorer_fails(tmp_path):
    config, storage = provision(SCHOOL_CONFIG, tmp_path)
    results = make_school_results(config, storage)

    result = results["s3"]
    assert result.overall_percentage == pytest.approx(30.27)
    assert result.grade == "F"
    assert result.passed is False

    failing = [c.course_name for c in result.course_results if not c.passed]
    assert failing == ["Mathematics", "Science", "English"]


def test_school_reports_generate_excel_and_pdf(tmp_path):
    config, storage = provision(SCHOOL_CONFIG, tmp_path)
    results = make_school_results(config, storage)

    output_dir = tmp_path / "reports"
    created = generate_reports(results["s1"], output_dir, config.reporting)

    names = [path.name for path in created]
    assert any(name.endswith(EXCEL_SUFFIX) for name in names)
    assert any(name.endswith(PDF_SUFFIX) for name in names)
    assert all(Path(path).exists() for path in created)
    assert sum(1 for path in created if path.suffix == ".png") == 1

    excel_path = next(path for path in created if path.suffix == ".xlsx")
    workbook = openpyxl.load_workbook(excel_path)
    sheet = workbook.active

    assert sheet["A1"].value == "Springfield Model School - Term 1 Result"
    assert sheet["B3"].value == "Aarav Sharma"
    assert sheet["B4"].value == "6A01"
    assert sheet["B5"].value == "Class 6 > Section A"
    assert sheet["B7"].value == "Term 1"

    # The main table contains the three courses.
    course_cells = [sheet.cell(row=10 + i, column=1).value for i in range(3)]
    assert course_cells == ["Mathematics", "Science", "English"]

    # One of the summary rows reports the overall percentage.
    labels = [sheet.cell(row=row, column=1).value for row in range(14, 18)]
    assert "Overall Percentage" in labels
    assert "Total Marks" in labels


# ------------------------------------------------------------
# College-style configuration
# ------------------------------------------------------------

COLLEGE_SCORES = {
    "c1": {
        "dsa_prog": {"Internal": 40, "External": 80},
        "dbms_prog": {"Internal": 30, "External": 55},
        "os_prog": {"Internal": 45, "External": 88},
    },
    "c2": {
        "dsa_prog": {"Internal": 20, "External": 25},
        "dbms_prog": {"Internal": 25, "External": 30},
        "os_prog": {"Internal": 28, "External": 32},
    },
}


def make_college_results(config, storage):
    by_offering = assessments_by_offering(storage, config.institution.id)
    record_scores(storage, config.institution.id, by_offering, COLLEGE_SCORES)
    context = make_context(config, storage)
    return {result.student_id: result for result in compute_all_results(context)}


def test_college_internal_external_weighted_result(tmp_path):
    config, storage = provision(COLLEGE_CONFIG, tmp_path)
    results = make_college_results(config, storage)

    result = results["c1"]
    assert result.student_name == "Meera Iyer"
    assert result.org_path == "Department CSE > Program B.Tech"
    assert result.term_name == "Semester 1"

    dsa = next(c for c in result.course_results if c.course_id == "dsa")
    assert dsa.percentage == pytest.approx(80)
    assert dsa.grade == "A"

    dbms = next(c for c in result.course_results if c.course_id == "dbms")
    assert dbms.percentage == pytest.approx(57)
    assert dbms.grade == "C"

    os_ = next(c for c in result.course_results if c.course_id == "os")
    assert os_.percentage == pytest.approx(88.8)
    assert os_.grade == "A"

    assert result.overall_percentage == pytest.approx(75.27)
    assert result.grade == "B+"
    assert result.passed is True


def test_college_low_scorer_fails(tmp_path):
    config, storage = provision(COLLEGE_CONFIG, tmp_path)
    results = make_college_results(config, storage)

    result = results["c2"]
    dsa = next(c for c in result.course_results if c.course_id == "dsa")
    assert dsa.percentage == pytest.approx(31)
    assert dsa.passed is False
    assert result.overall_percentage == pytest.approx(36.87)
    assert result.grade == "F"
    assert result.passed is False


def test_college_reports_generate(tmp_path):
    config, storage = provision(COLLEGE_CONFIG, tmp_path)
    results = make_college_results(config, storage)

    output_dir = tmp_path / "college_reports"
    created = generate_reports(results["c1"], output_dir, config.reporting)
    assert any(name.endswith(EXCEL_SUFFIX) for name in [path.name for path in created])
    assert any(name.endswith(PDF_SUFFIX) for name in [path.name for path in created])
    assert all(Path(path).exists() for path in created)