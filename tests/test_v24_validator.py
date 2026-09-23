"""Tests for V2.4 Phase 4: the Intelligent Validator.

The validator consumes the Phase 3 :class:`WorkbookNormalization` (and,
optionally, the Phase 2 detection) and produces a structured
:class:`ValidationResult` that separates errors, warnings and
review-required conditions.  It validates evidence only, never invents
missing information (no invented totals, no invented grading scale, no
assumed assessment scheme), preserves the normalizer's provenance, and
keeps ``review_required`` distinct from ``error`` so ambiguity is never
auto-promoted to invalidity.  These tests exercise the rules directly on
constructed records and end-to-end against the real workspace workbooks.
"""

import pytest
from openpyxl import Workbook

from assessment_engine.config import load_config
from assessment_engine.intelligent_import import (
    CellLocation,
    NormalizedAssessment,
    NormalizedField,
    NormalizedRecord,
    NormalizedSheet,
    NormalizedValue,
    SEVERITY_ERROR,
    SEVERITY_REVIEW,
    SEVERITY_WARNING,
    ValidationIssue,
    VALUE_GRADE,
    VALUE_IDENTIFIER,
    VALUE_MARK,
    VALUE_PERCENTAGE,
    VALUE_TOTAL,
    VALUE_UNRESOLVED,
    VALUE_ZERO,
    WorkbookNormalization,
    build_validation_context,
    normalize_workbook,
    validate_record,
    validate_workbook,
    validate_workbook_file,
)
from assessment_engine.intelligent_import.validator.rules import (
    RULE_AMB_UNRESOLVED,
    RULE_ASS_EXCEEDS_MAX,
    RULE_ASS_WITHOUT_SUBJECT,
    RULE_DUP_CONFLICTING,
    RULE_DUP_EXACT,
    RULE_GRD_NO_SCALE,
    RULE_GRD_SCALE_MISMATCH,
    RULE_MARK_EXCEEDS_TOTAL,
    RULE_MARK_NEGATIVE,
    RULE_MISS_IDENTITY,
    RULE_STR_CONFLICTING_ASSESSMENTS,
    RULE_STR_MARK_WITHOUT_SUBJECT,
    RULE_TOT_ASSESSMENT_SUM,
    RULE_TOT_NEGATIVE,
)

config_path = "configs/college.json"


# ------------------------------------------------------------
# Helpers - construct normalized records directly
# ------------------------------------------------------------


def _loc(sheet="Results", row=2, column=2, header="Marks"):
    return CellLocation(sheet, row, column, [], header)


def _value(kind, normalized, raw=None, row=2, column=2, header="Marks", extra=None):
    return NormalizedValue(
        kind, raw, normalized, location=_loc(row=row, column=column, header=header),
        extra=extra or {},
    )


def make_record(
    sheet="Results",
    row=2,
    identity=None,
    subject=None,
    subject_mark=None,
    percentage=None,
    total=None,
    grade=None,
    assessments=None,
    unresolved=None,
    statuses=None,
    measures=None,
):
    """Build a NormalizedRecord for validator unit tests."""
    record = NormalizedRecord(sheet=sheet, row=row)
    for kind, value in (identity or []):
        record.identity.append(NormalizedField(kind, value))
    if subject is not None:
        record.subject = NormalizedField(
            "subject",
            NormalizedValue(
                VALUE_IDENTIFIER if False else "text",
                subject,
                subject,
                location=_loc(row=row, column=1, header="Subject"),
            ),
        )
    if subject_mark is not None:
        record.subject_mark = NormalizedField(
            "subject_mark",
            _value(VALUE_MARK, float(subject_mark), raw=subject_mark,
                   row=row, column=2, header="Marks"),
        )
    if percentage is not None:
        record.measures.append(
            NormalizedField(
                "percentage",
                _value(VALUE_PERCENTAGE, float(percentage), raw=percentage,
                       row=row, column=3, header="Percentage"),
            )
        )
    if total is not None:
        if isinstance(total, tuple):
            raw, maximum = total
            record.measures.append(
                NormalizedField(
                    "total_marks",
                    _value(VALUE_TOTAL, float(raw), raw=f"{raw}/{maximum}",
                           row=row, column=4, header="Total Marks",
                           extra={"maximum": float(maximum)}),
                )
            )
        else:
            record.measures.append(
                NormalizedField(
                    "total_marks",
                    _value(VALUE_TOTAL, float(total), raw=total,
                           row=row, column=4, header="Total Marks"),
                )
            )
    if grade is not None:
        record.measures.append(
            NormalizedField(
                "grade",
                _value(VALUE_GRADE, grade, raw=grade,
                       row=row, column=5, header="Grade"),
            )
        )
    for name, mark in (assessments or []):
        record.assessments.append(
            NormalizedAssessment(
                name=NormalizedValue(
                    "text", name, name,
                    location=_loc(row=row, column=6, header=name),
                ),
                mark=NormalizedValue(
                    VALUE_MARK, mark, float(mark),
                    location=_loc(row=row, column=6, header=name),
                ),
            )
        )
    for field_kind, value in (unresolved or []):
        record.unresolved.append(
            NormalizedField(field_kind, value)
        )
    for field_kind, value in (statuses or []):
        record.statuses.append(NormalizedField(field_kind, value))
    for field_kind, value in (measures or []):
        record.measures.append(NormalizedField(field_kind, value))
    return record


def make_workbook_normalization(records, source="unit.xlsx", filename="unit.xlsx",
                                warnings=None, sheets=None):
    if sheets is None:
        by_sheet = {}
        for record in records:
            by_sheet.setdefault(record.sheet, []).append(record)
        sheets = [
            NormalizedSheet(name=name, records=items)
            for name, items in by_sheet.items()
        ]
    return WorkbookNormalization(
        source=source,
        filename=filename,
        sheets=sheets,
        records=records,
        warnings=list(warnings or []),
    )


def roll_no(raw, row=2):
    return ("roll_no", NormalizedValue(
        VALUE_IDENTIFIER, raw, str(raw),
        location=_loc(row=row, column=0, header="Roll No"),
    ))


def no_context():
    return build_validation_context()


def college_context():
    return build_validation_context(config=load_config(config_path))


# ------------------------------------------------------------
# 1. Record structure
# ------------------------------------------------------------


def test_valid_subject_mark_percentage_grade_record():
    record = make_record(
        subject="AI", subject_mark=69, percentage=69, grade="B",
    )
    validation = validate_record(record, no_context())
    assert validation.accepted is True
    assert validation.error_count == 0
    # grade unverifiable without a scale, identity absent: both reviews, not errors
    assert validation.issues_for_rule(RULE_GRD_NO_SCALE)
    assert all(issue.severity == SEVERITY_REVIEW for issue in validation.issues)


def test_subject_mark_without_subject_is_review():
    record = make_record(subject_mark=69, percentage=69)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_STR_MARK_WITHOUT_SUBJECT)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW


def test_contradictory_assessment_subjects_are_review():
    record = make_record(subject="Physics", subject_mark=20)
    record.assessments.append(
        NormalizedAssessment(
            name=NormalizedValue("text", "SEE", "see", location=_loc()),
            mark=NormalizedValue(VALUE_MARK, 50, 50.0, location=_loc()),
            subject=NormalizedValue("text", "Maths", "maths", location=_loc()),
        )
    )
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_STR_CONFLICTING_ASSESSMENTS)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW


# ------------------------------------------------------------
# 2. Numeric values
# ------------------------------------------------------------


def test_negative_mark_is_error():
    record = make_record(subject="AI", subject_mark=-5)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_MARK_NEGATIVE)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert validation.accepted is False


def test_negative_assessment_mark_is_error():
    record = make_record(
        subject="Physics",
        assessments=[("Internal", -3), ("External", 50)],
    )
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_MARK_NEGATIVE)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR


def test_percentage_above_hundred_is_error():
    record = make_record(subject="AI", percentage=120)
    validation = validate_record(record, no_context())
    pct_high = validation.issues_for_rule("V24-PCT-002")
    assert len(pct_high) == 1
    assert pct_high[0].severity == SEVERITY_ERROR


def test_percentage_below_zero_is_error():
    record = make_record(subject="AI", percentage=-5)
    validation = validate_record(record, no_context())
    pct_low = validation.issues_for_rule("V24-PCT-001")
    assert len(pct_low) == 1
    assert pct_low[0].severity == SEVERITY_ERROR


def test_obtained_marks_above_total_is_error():
    record = make_record(subject="AI", subject_mark=90, total=80)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_MARK_EXCEEDS_TOTAL)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert validation.accepted is False


def test_negative_total_is_error():
    record = make_record(subject="AI", subject_mark=10, total=-20)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_TOT_NEGATIVE)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR


# ------------------------------------------------------------
# 3. Percentage consistency
# ------------------------------------------------------------


def test_percentage_mismatch_when_total_known():
    record = make_record(subject="AI", subject_mark=75, total=100, percentage=80)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule("V24-PCT-003")
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert issue[0].evidence["marks"] == 75.0
    assert issue[0].evidence["maximum"] == 100.0


def test_missing_total_with_percentage_does_not_invent_total():
    record = make_record(subject="AI", subject_mark=75, percentage=80)
    validation = validate_record(record, no_context())
    assert validation.issues_for_rule("V24-PCT-003") == []
    assert validation.error_count == 0


def test_percentage_consistent_with_fraction_total():
    record = make_record(
        subject="AI", subject_mark=69, total=(69, 100), percentage=69,
    )
    validation = validate_record(record, no_context())
    assert validation.issues_for_rule("V24-PCT-003") == []
    assert validation.error_count == 0


def test_does_not_assume_marks_out_of_100():
    record = make_record(subject="AI", subject_mark=80, percentage=100)
    assert validate_record(record, no_context()).error_count == 0


# ------------------------------------------------------------
# 4. Grade consistency
# ------------------------------------------------------------


def test_grade_without_scale_is_review_not_error():
    record = make_record(subject="AI", subject_mark=69, percentage=69, grade="B")
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_GRD_NO_SCALE)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW
    assert validation.error_count == 0


def test_grade_satisfies_configured_scale():
    record = make_record(subject="AI", subject_mark=69, percentage=69, grade="B")
    validation = validate_record(record, college_context())
    assert validation.issues_for_rule(RULE_GRD_SCALE_MISMATCH) == []
    assert validation.issues_for_rule(RULE_GRD_NO_SCALE) == []


def test_grade_mismatch_with_configured_scale_is_warning():
    record = make_record(subject="AI", subject_mark=92, percentage=92, grade="C")
    validation = validate_record(record, college_context())
    issue = validation.issues_for_rule(RULE_GRD_SCALE_MISMATCH)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_WARNING
    assert not validation.issues_for_rule(RULE_GRD_NO_SCALE)


# ------------------------------------------------------------
# 5. Assessment relationships
# ------------------------------------------------------------


def test_assessment_mark_above_configured_maximum_is_error():
    record = make_record(
        identity=[roll_no("1")],
        assessments=[("Internal", 60), ("External", 90)],
    )
    validation = validate_record(record, college_context())
    issue = validation.issues_for_rule(RULE_ASS_EXCEEDS_MAX)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert issue[0].evidence == {"assessment": "Internal", "mark": 60.0, "maximum": 50.0}


def test_assessment_without_subject_is_review():
    record = make_record(
        identity=[roll_no("1")],
        assessments=[("Internal", 20), ("External", 60)],
    )
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_ASS_WITHOUT_SUBJECT)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW
    assert validation.error_count == 0


def test_assessment_sum_above_total_is_error():
    record = make_record(
        identity=[roll_no("1")],
        total=70,
        assessments=[("Internal", 50), ("External", 60)],
    )
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_TOT_ASSESSMENT_SUM)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert issue[0].evidence["sum"] == 110.0


def test_assessment_sum_within_total_is_fine():
    record = make_record(
        identity=[roll_no("1")],
        total=80,
        assessments=[("Internal", 20), ("External", 60)],
    )
    validation = validate_record(record, no_context())
    assert validation.issues_for_rule(RULE_TOT_ASSESSMENT_SUM) == []
    assert validation.error_count == 0


# ------------------------------------------------------------
# 6. Missing data (acceptable vs review)
# ------------------------------------------------------------


def test_missing_total_is_acceptable():
    record = make_record(subject="AI", subject_mark=69, percentage=69)
    validation = validate_record(record, no_context())
    # no total is present, so no total-derived inconsistency is invented
    assert validation.issues_for_rule("V24-PCT-003") == []
    assert validation.issues_for_rule(RULE_MARK_EXCEEDS_TOTAL) == []
    # identity is the only review concern here
    assert [issue.rule_id for issue in validation.issues] == [RULE_MISS_IDENTITY]
    assert validation.error_count == 0


def test_missing_identity_is_review_concern():
    record = make_record(subject="AI", subject_mark=69, percentage=69)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_MISS_IDENTITY)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW
    assert "not invented" in issue[0].message


def test_identity_present_avoids_missing_identity_review():
    record = make_record(
        identity=[roll_no("2024-001")],
        subject="AI", subject_mark=69, percentage=69,
    )
    validation = validate_record(record, no_context())
    assert validation.issues_for_rule(RULE_MISS_IDENTITY) == []


def test_preserved_ambiguity_is_review_not_error():
    value = NormalizedValue(
        VALUE_UNRESOLVED, 75, None,
        location=_loc(), note="column meaning is unresolved (ambiguous_measure)",
    )
    record = make_record(unresolved=[("marks", value)])
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_AMB_UNRESOLVED)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW
    assert validation.error_count == 0


# ------------------------------------------------------------
# 7. Provenance
# ------------------------------------------------------------


def test_validation_issue_provenance():
    record = make_record(subject="AI", subject_mark=-5, row=7)
    validation = validate_record(record, no_context())
    issue = validation.issues_for_rule(RULE_MARK_NEGATIVE)[0]
    assert isinstance(issue, ValidationIssue)
    assert issue.sheet == "Results"
    assert issue.row == 7
    assert issue.column == 2
    assert issue.header_text == "Marks"
    assert issue.field == "subject_mark"
    assert issue.location.render() == "Results!C7"


def test_issue_render():
    record = make_record(subject="AI", subject_mark=-5, row=3)
    issue = validate_record(record, no_context()).issues_for_rule(RULE_MARK_NEGATIVE)[0]
    rendered = issue.render()
    assert "V24-MARK-001" in rendered
    assert "error" in rendered
    assert "Results!C3" in rendered


# ------------------------------------------------------------
# 8. Workbook-level rules: duplicates and conflicts
# ------------------------------------------------------------


def test_exact_duplicate_records_are_warning():
    first = make_record(
        sheet="Results", row=2,
        identity=[roll_no("1")], subject="AI", subject_mark=69, percentage=69,
    )
    second = make_record(
        sheet="Results", row=3,
        identity=[roll_no("1")], subject="AI", subject_mark=69, percentage=69,
    )
    result = validate_workbook(make_workbook_normalization([first, second]), context=no_context())
    issue = result.issues_for_rule(RULE_DUP_EXACT)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_WARNING
    assert issue[0].evidence["original_row"] == 2
    assert issue[0].evidence["duplicate_row"] == 3
    assert result.valid is True  # duplicates alone do not invalidate


def test_conflicting_duplicate_records_are_review():
    first = make_record(
        sheet="Results", row=2,
        identity=[roll_no("1")], subject="AI", subject_mark=69, percentage=69,
    )
    second = make_record(
        sheet="Results", row=3,
        identity=[roll_no("1")], subject="AI", subject_mark=88, percentage=88,
    )
    result = validate_workbook(make_workbook_normalization([first, second]), context=no_context())
    issue = result.issues_for_rule(RULE_DUP_CONFLICTING)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_REVIEW
    assert result.valid is True  # conflicting evidence is review, not error


def test_duplicate_requires_subject_evidence():
    first = make_record(sheet="Results", row=2, identity=[roll_no("1")], subject_mark=69)
    second = make_record(sheet="Results", row=3, identity=[roll_no("1")], subject_mark=69)
    result = validate_workbook(make_workbook_normalization([first, second]), context=no_context())
    assert result.issues_for_rule(RULE_DUP_EXACT) == []
    assert result.issues_for_rule(RULE_DUP_CONFLICTING) == []


def test_insufficient_identity_evidence_never_claims_duplicates():
    first = make_record(
        sheet="Results", row=2, subject="AI", subject_mark=69,
        identity=[("serial_no", NormalizedValue(VALUE_IDENTIFIER, 1, "1", location=_loc(column=0)))] ,
    )
    second = make_record(
        sheet="Results", row=3, subject="AI", subject_mark=69,
        identity=[("serial_no", NormalizedValue(VALUE_IDENTIFIER, "2", "2", location=_loc(column=0)))],
    )
    result = validate_workbook(make_workbook_normalization([first, second]), context=no_context())
    assert result.issues_for_rule(RULE_DUP_EXACT) == []
    assert result.issues_for_rule(RULE_DUP_CONFLICTING) == []


def test_serial_number_is_not_student_identity():
    record = make_record(
        subject="AI", subject_mark=69,
        identity=[("serial_no", NormalizedValue(VALUE_IDENTIFIER, 1, "1", location=_loc(column=0)))],
    )
    validation = validate_record(record, no_context())
    assert validation.issues_for_rule(RULE_MISS_IDENTITY)


def test_distinct_assessments_are_not_duplicates():
    ia = make_record(
        sheet="Results", row=2, identity=[roll_no("1")], subject="AI",
        assessments=[("IA", 20)],
    )
    see = make_record(
        sheet="Results", row=3, identity=[roll_no("1")], subject="AI",
        assessments=[("SEE", 50)],
    )
    result = validate_workbook(make_workbook_normalization([ia, see]), context=no_context())
    assert result.issues_for_rule(RULE_DUP_EXACT) == []
    assert result.issues_for_rule(RULE_DUP_CONFLICTING) == []


def test_same_assessment_duplicated_once_is_warning():
    ia1 = make_record(
        sheet="Results", row=2, identity=[roll_no("1")], subject="AI",
        assessments=[("IA", 20)],
    )
    ia2 = make_record(
        sheet="Results", row=3, identity=[roll_no("1")], subject="AI",
        assessments=[("IA", 20)],
    )
    result = validate_workbook(make_workbook_normalization([ia1, ia2]), context=no_context())
    assert len(result.issues_for_rule(RULE_DUP_EXACT)) == 1


# ------------------------------------------------------------
# 9. Severity separation and result semantics
# ------------------------------------------------------------


def test_warning_vs_error_vs_review_separation():
    clean = make_record(
        identity=[roll_no("1")], subject="AI", subject_mark=69, percentage=69,
    )
    too_high_mark = make_record(
        identity=[roll_no("2")], subject="Maths", subject_mark=95, percentage=95, total=80,
    )
    conflicting_pct = make_record(
        identity=[roll_no("3")], subject="Physics", subject_mark=70, percentage=55, total=100,
    )
    result = validate_workbook(
        make_workbook_normalization([clean, too_high_mark, conflicting_pct]),
        context=college_context(),
    )
    assert result.valid is False
    # mark 95/80 -> V24-MARK-002 + V24-PCT-003 (95% can never match 118.75%);
    # mark 70/100 vs 55% -> V24-PCT-003. Record 1 stays error-free.
    assert result.error_count == 3
    assert result.records[0].error_count == 0
    assert result.records[1].error_count == 2
    assert result.records[2].error_count == 1
    assert result.records[0].accepted is True
    assert result.records[1].accepted is False
    assert len(result.accepted_records) == 1
    assert len(result.rejected_records) == 2


def test_result_is_serializable_dict():
    record = make_record(subject="AI", subject_mark=-5, row=7)
    result = validate_workbook(make_workbook_normalization([record]), context=no_context())
    payload = result.as_dict()
    assert payload["valid"] is False
    assert payload["errors"] == 1
    assert payload["records"][0]["accepted"] is False
    assert payload["issues"][0]["location"] == "Results!C7"


# ------------------------------------------------------------
# 10. End-to-end pipeline integration
# ------------------------------------------------------------


def write_workbook(path, sheets):
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


def test_pipeline_subject_result_table_valid(tmp_path):
    path = write_workbook(
        tmp_path / "subject_result.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                ["1", "AI", 69, 0.69, "B"],
                ["2", "Mathematics", 94, 0.94, "A+"],
            ]
        },
    )
    normalization = normalize_workbook(path)
    result = validate_workbook(normalization, context=college_context())
    assert result.valid is True
    assert result.filename == "subject_result.xlsx"
    assert len(result.records) == 2
    # marks, percentages and grades all live inside the accepted records
    assert all(item.accepted for item in result.records)
    # missing student identity is a review concern, not an error
    assert result.issues_for_rule(RULE_MISS_IDENTITY)
    assert result.issues_for_rule(RULE_GRD_SCALE_MISMATCH) == []


def test_pipeline_assessment_relationships(tmp_path):
    path = write_workbook(
        tmp_path / "assessment.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["1", "Aarav", 20, 60],
            ]
        },
    )
    result = validate_workbook_file(path, config=config_path)
    assert result.valid is True
    # assessments carry no subject relationship: review-worthy, not an error
    assert result.issues_for_rule(RULE_ASS_WITHOUT_SUBJECT)
    assert result.issues_for_rule(RULE_ASS_EXCEEDS_MAX) == []
    assert result.error_count == 0


def test_pipeline_assessment_over_configured_max(tmp_path):
    path = write_workbook(
        tmp_path / "assessment_over.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["1", "Aarav", 60, 60],
            ]
        },
    )
    result = validate_workbook_file(path, config=config_path)
    issue = result.issues_for_rule(RULE_ASS_EXCEEDS_MAX)
    assert len(issue) == 1
    assert issue[0].severity == SEVERITY_ERROR
    assert result.valid is False


def test_pipeline_negative_mark_invalidates(tmp_path):
    path = write_workbook(
        tmp_path / "negative.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                ["1", "AI", -5, 0.05, "B"],
            ]
        },
    )
    result = validate_workbook_file(path)
    assert result.issues_for_rule(RULE_MARK_NEGATIVE)
    assert result.valid is False


# ------------------------------------------------------------
# 11. Real semester-result workbooks
# ------------------------------------------------------------


@pytest.fixture(scope="module")
def real_workbooks():
    import pathlib

    base = pathlib.Path("Semester_Result")
    names = [
        "Quaja_Hussain_Semester_Result.xlsx",
        "Priya_Sharma_Semester_Result.xlsx",
        "MS_Mohammed_Quaja_Hussain_Semester_Result.xlsx",
    ]
    return [
        base / name for name in names if (base / name).is_file()
    ]


def test_real_workbook_pipeline_validation(real_workbooks):
    if not real_workbooks:
        pytest.skip("no real semester-result workbooks present")
    for path in real_workbooks:
        result = validate_workbook_file(
            path, config=config_path
        )
        assert result.valid is True, path
        assert result.error_count == 0, path
        assert len(result.records) >= 1, path
        assert all(item.accepted for item in result.records), path
        # identity is absent from these sheets, so a review concern is expected
        assert result.issues_for_rule(RULE_MISS_IDENTITY), path
        # configured scale grades match the recorded grades
        assert result.issues_for_rule(RULE_GRD_SCALE_MISMATCH) == [], path


def test_real_workbook_public_api(real_workbooks):
    if not real_workbooks:
        pytest.skip("no real semester-result workbooks present")
    path = real_workbooks[0]
    from assessment_engine.intelligent_import import (
        detect_workbook,
        inspect_workbook,
        normalize_workbook_profile,
    )

    profile = inspect_workbook(path)
    detection = detect_workbook(profile)
    normalization = normalize_workbook_profile(profile, detection=detection)
    result = validate_workbook(normalization, detection=detection, context=college_context())
    assert result.valid is True
    assert result.filename == path.name
    for issue in result.issues:
        if issue.rule_id == RULE_MISS_IDENTITY:
            assert issue.sheet == "Semester Result"
            assert issue.row is not None