"""V2.3 Step 1: multi-tenancy foundation.

Verifies the Institution-as-tenant boundary: a tenant registry
(``Storage.list_institutions``) plus cross-tenant isolation for the
existing JSON storage.  No V2.2 behaviour changes.
"""

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from assessment_engine import importer, models
from assessment_engine.analyzer import provision_institution
from assessment_engine.cli import cmd_import, cmd_init, cmd_list, cmd_report
from assessment_engine.config import load_config
from assessment_engine.storage import JsonStorage, Storage

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHOOL_CONFIG = REPO_ROOT / "configs" / "school.json"
COLLEGE_CONFIG = REPO_ROOT / "configs" / "college.json"

SCHOOL_ASSESSMENTS = ["FA-1", "FA-2", "SA-1", "SA-2"]


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def make_two_tenant_storage(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    provision_institution(load_config(SCHOOL_CONFIG), storage)
    provision_institution(load_config(COLLEGE_CONFIG), storage)
    return storage


def maths_headers():
    return ["Roll No", "Student Name"] + SCHOOL_ASSESSMENTS


def write_workbook(path, sheets):
    """sheets: dict {sheet_name: list_of_rows}."""
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


# ------------------------------------------------------------
# Tenant registry
# ------------------------------------------------------------


def test_empty_data_dir_returns_empty_list(tmp_path):
    storage = JsonStorage(tmp_path / "data")
    assert storage.list_institutions() == []


def test_two_institutions_return_sorted_ids(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    assert storage.list_institutions() == ["sample_college", "sample_school"]


def test_unrelated_dirs_and_files_are_ignored(tmp_path):
    data_dir = tmp_path / "data"
    (data_dir / "sample_school").mkdir(parents=True)
    (data_dir / "sample_school" / "db.json").write_text(
        '{"institution": {"id": "sample_school", "name": "S"}}',
        encoding="utf-8",
    )
    (data_dir / "scratch").mkdir()
    (data_dir / "scratch" / "notes.txt").write_text("not a tenant", encoding="utf-8")
    (data_dir / "stray.txt").write_text("not a tenant", encoding="utf-8")

    storage = JsonStorage(data_dir)
    assert storage.list_institutions() == ["sample_school"]


# ------------------------------------------------------------
# Cross-tenant isolation
# ------------------------------------------------------------


def test_cross_tenant_student_isolation(tmp_path):
    storage = make_two_tenant_storage(tmp_path)

    school_students = storage.load_all(models.Student, "sample_school")
    college_students = storage.load_all(models.Student, "sample_college")

    assert {s.name for s in school_students} == {
        "Aarav Sharma",
        "Bhavna Patel",
        "Rohan Mehta",
    }
    assert {s.name for s in college_students} == {"Meera Iyer", "Karthik Rao"}


def test_v22_db_json_stays_readable_and_discoverable(tmp_path):
    data_dir = tmp_path / "data"
    db_path = data_dir / "legacy_inst" / "db.json"
    db_path.parent.mkdir(parents=True)
    db_path.write_text(
        json.dumps(
            {
                "institution": {
                    "id": "legacy_inst",
                    "name": "Legacy School",
                    "org_levels": ["Class", "Section"],
                },
                "student": {
                    "s1": {
                        "id": "s1",
                        "institution_id": "legacy_inst",
                        "name": "Ada",
                        "roll_no": "1",
                        "extra": {},
                    }
                },
            }
        ),
        encoding="utf-8",
    )

    storage = JsonStorage(data_dir)
    assert storage.list_institutions() == ["legacy_inst"]

    institution = storage.load(models.Institution, "legacy_inst", "legacy_inst")
    assert institution is not None
    assert institution.name == "Legacy School"

    students = storage.load_all(models.Student, "legacy_inst")
    assert [student.name for student in students] == ["Ada"]


def test_custom_storage_subclass_stays_instantiable():
    class LegacyStorage(Storage):
        def save(self, record):
            return None

        def load(self, model, obj_id, institution_id):
            return None

        def load_all(self, model, institution_id):
            return []

        def clear(self, institution_id):
            return None

    storage = LegacyStorage()
    assert isinstance(storage, Storage)
    with pytest.raises(NotImplementedError):
        storage.list_institutions()


def test_import_for_tenant_a_does_not_modify_tenant_b(tmp_path):
    storage = make_two_tenant_storage(tmp_path)
    config = load_config(SCHOOL_CONFIG)

    before_a_marks = len(storage.load_all(models.Mark, "sample_school"))
    before_b_marks = len(storage.load_all(models.Mark, "sample_college"))
    before_b_students = len(storage.load_all(models.Student, "sample_college"))

    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers(), ["6A01", "Aarav Sharma", 20, 22, 78, 85]]},
    )
    plan = importer.parse_workbook(config, storage, workbook_path)
    assert plan.has_errors is False
    plan.apply(storage)

    assert len(storage.load_all(models.Mark, "sample_school")) == before_a_marks + 4
    assert len(storage.load_all(models.Mark, "sample_college")) == before_b_marks
    assert len(storage.load_all(models.Student, "sample_college")) == before_b_students


# ------------------------------------------------------------
# CLI
# ------------------------------------------------------------


def test_cli_list_command(tmp_path, capsys):
    storage = make_two_tenant_storage(tmp_path)

    class Args:
        data = str(storage.data_dir)

    assert cmd_list(Args()) == 0
    captured = capsys.readouterr().out.splitlines()
    assert captured == ["sample_college", "sample_school"]


def test_existing_init_import_report_behavior_unchanged(tmp_path):
    data_dir = str(tmp_path / "data")

    class InitArgs:
        config_path = str(SCHOOL_CONFIG)
        data = data_dir

    assert cmd_init(InitArgs()) == 0

    workbook_path = write_workbook(
        tmp_path / "marks.xlsx",
        {"maths_a": [maths_headers(), ["6A01", "Aarav Sharma", 20, 22, 78, 85]]},
    )

    class ImportArgs:
        config_path = str(SCHOOL_CONFIG)
        data = data_dir
        file = str(workbook_path)
        template = None
        offering = None
        legacy = False
        skip_invalid = False
        yes = True

    assert cmd_import(ImportArgs()) == 0

    class ReportArgs:
        config_path = str(SCHOOL_CONFIG)
        data = data_dir
        student = "s1"

    assert cmd_report(ReportArgs()) == 0

    storage = JsonStorage(data_dir)
    assert storage.list_institutions() == ["sample_school"]
    assert len(storage.load_all(models.Mark, "sample_school")) == 4

    reports = storage.data_dir / "sample_school" / "reports"
    assert list(reports.glob("*.xlsx"))
    assert list(reports.glob("*.pdf"))