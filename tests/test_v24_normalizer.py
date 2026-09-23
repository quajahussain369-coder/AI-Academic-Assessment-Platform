"""Tests for V2.4 Phase 3: the Intelligent Normalizer.

The normalizer turns a Phase 1 structural profile plus a Phase 2 semantic
detection into a canonical, deterministic representation.  Every
normalized value must keep its raw source value and full provenance
(sheet, row, column and header path).  Marks, percentages, grades, totals
and academic status tokens (AB/Absent/Exempt/zero/blank) must be
normalized deterministically, and phase-2 ambiguities must be preserved -
never guessed.  These tests verify the canonical output (and that the
normalizer never touches ImportPlan/importing behaviour).
"""

import pytest
from openpyxl import Workbook

from assessment_engine.intelligent_import import (
    CellLocation,
    NormalizedRecord,
    NormalizedValue,
    VALUE_BLANK,
    VALUE_GRADE,
    VALUE_IDENTIFIER,
    VALUE_MARK,
    VALUE_PERCENTAGE,
    VALUE_STATUS,
    VALUE_UNRESOLVED,
    VALUE_ZERO,
    WorkbookNormalization,
    detect_workbook,
    inspect_workbook,
    inspect_sheet,
    normalize_sheet,
    normalize_workbook,
    normalize_workbook_profile,
)
from assessment_engine.intelligent_import.normalizer.value import (
    normalize_grade,
    normalize_identifier,
    normalize_mark,
    normalize_percentage,
    normalize_status,
    normalize_text,
    normalize_total,
)


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def write_workbook(path, sheets):
    """sheets: dict {sheet_name: list_of_rows}."""
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


def normalize(path):
    """Full pipeline: inspect -> detect -> normalize."""
    return normalize_workbook(path)


# ------------------------------------------------------------
# Value helpers
# ------------------------------------------------------------


def test_normalize_mark_preserves_status_tokens_and_zeros():
    assert normalize_mark("AB").kind == VALUE_STATUS
    assert normalize_mark("AB").normalized == "absent"
    assert normalize_mark("Absent").normalized == "absent"
    assert normalize_mark("Exempt").normalized == "exempt"
    assert normalize_mark("EX").normalized == "exempt"
    assert normalize_mark(0).kind == VALUE_ZERO
    assert normalize_mark(0).normalized == 0.0
    assert normalize_mark(None).kind == VALUE_BLANK
    assert normalize_mark("").kind == VALUE_BLANK
    assert normalize_mark("-").kind == VALUE_BLANK
    assert normalize_mark("69").kind == VALUE_MARK
    assert normalize_mark("69").normalized == 69.0
    assert normalize_mark("mystery").kind == VALUE_UNRESOLVED
    assert normalize_mark("AB").raw == "AB"


def test_normalize_percentage_distinguishes_forms():
    assert normalize_percentage("69%").normalized == 69.0
    assert normalize_percentage("69%").note  # explicit % suffix
    assert normalize_percentage(0.69).normalized == 69.0
    assert normalize_percentage(0.69).raw == 0.69
    assert normalize_percentage(69).normalized == 69.0
    assert normalize_percentage("69").normalized == 69.0
    assert normalize_percentage(0.0).kind == VALUE_ZERO
    assert normalize_percentage("not a number").kind == VALUE_UNRESOLVED
    # raw values stay distinguishable
    forms = {normalize_percentage(v).raw: normalize_percentage(v).normalized
             for v in ("69%", 0.69, 69)}
    assert forms[0.69] == forms[69] == 69.0


def test_normalize_total_fraction_keeps_maximum():
    value = normalize_total("69/100")
    assert value.kind == "total"
    assert value.normalized == 69.0
    assert value.extra == {"maximum": 100.0}
    assert value.raw == "69/100"
    assert normalize_total(80).normalized == 80.0
    assert normalize_total("not a number").kind == VALUE_UNRESOLVED


def test_normalize_grade_and_identifier_and_text():
    assert normalize_grade("A+").kind == VALUE_GRADE
    assert normalize_grade("A+").normalized == "A+"
    assert normalize_grade("   b ").normalized == "b"
    assert normalize_identifier("CSE01").kind == VALUE_IDENTIFIER
    assert normalize_identifier(101).normalized == "101"
    assert normalize_text(" Data Structures ").normalized == "Data Structures"


def test_normalize_status_only_known_tokens():
    assert normalize_status("AB") == "absent"
    assert normalize_status("abs") == "absent"
    assert normalize_status("Exempt") == "exempt"
    assert normalize_status("Pass") is None
    assert normalize_status("") is None
    assert normalize_status(None) is None


def test_cell_location_render():
    assert CellLocation("Semester Result", 9, 2).render() == "Semester Result!C9"
    assert CellLocation("Results", 2, 4).render() == "Results!E2"
    assert CellLocation("A", 2, 0).column_letter == "A"
    assert CellLocation("A", 2, 25).column_letter == "Z"
    assert CellLocation("A", 2, 26).column_letter == "AA"
    assert CellLocation("A", 9).render() == "A row 9"


# ------------------------------------------------------------
# 1. Subject result table -> subject-level records
# ------------------------------------------------------------


def test_subject_result_table_produces_subject_records(tmp_path):
    path = write_workbook(
        tmp_path / "subject_result.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                ["1", "Physics", 76, 0.76, "A"],
                ["2", "Mathematics", 82, 0.82, "A+"],
            ]
        },
    )
    normalization = normalize(path)
    sheet = normalization.sheet("Results")
    assert len(sheet.records) == 2
    assert sheet.unresolved == []

    record = sheet.records[0]
    assert isinstance(record, NormalizedRecord)
    assert record.subject is not None
    assert record.subject.kind == "subject"
    assert record.subject.value.normalized == "Physics"

    # subject-level rather than assessment records
    assert record.assessments == []
    assert record.subject_mark is not None
    assert record.subject_mark.value.kind == VALUE_MARK
    assert record.subject_mark.value.normalized == 76.0
    assert record.subject_mark.value.raw == 76

    assert record.find("percentage").value.kind == VALUE_PERCENTAGE
    assert record.find("percentage").value.normalized == 76.0
    assert record.find("percentage").value.raw == 0.76
    assert record.find("grade").value.kind == VALUE_GRADE
    assert record.find("grade").value.normalized == "A"

    second = sheet.records[1]
    assert second.subject.value.normalized == "Mathematics"
    assert second.subject_mark.value.normalized == 82.0
    assert second.find("percentage").value.normalized == 82.0
    assert second.find("grade").value.normalized == "A+"


# ------------------------------------------------------------
# 2. Assessment-oriented table -> assessment records
# ------------------------------------------------------------


def test_assessment_table_produces_assessment_records(tmp_path):
    path = write_workbook(
        tmp_path / "assessment_table.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["1", "Aarav", 20, 60],
                ["2", "Bhavna", 22, 58],
            ]
        },
    )
    sheet = normalize(path).sheet("Results")
    assert len(sheet.records) == 2

    record = sheet.records[0]
    assert record.find("roll_no").value.kind == VALUE_IDENTIFIER
    assert record.find("roll_no").value.normalized == "1"
    assert record.find("student_name").value.normalized == "Aarav"

    assert len(record.assessments) == 2
    internal, external = record.assessments
    assert internal.name.raw == "Internal"
    assert internal.name.normalized == "internal"
    assert internal.mark.kind == VALUE_MARK
    assert internal.mark.normalized == 20.0
    assert external.name.normalized == "external"
    assert external.mark.normalized == 60.0

    assert record.subject is None
    assert record.measures == []
    assert record.unresolved == []


# ------------------------------------------------------------
# 3. Multi-row header hierarchy -> Physics/IA, Physics/SEE
# ------------------------------------------------------------


def test_multi_row_header_preserves_subject_assessment(tmp_path):
    path = write_workbook(
        tmp_path / "multi.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Physics", "Physics", "Maths", "Maths"],
                ["", "", "IA", "SEE", "IA", "SEE"],
                ["1", "Aarav", 10, 50, 12, 40],
                ["2", "Bhavna", 11, 52, 13, 41],
            ]
        },
    )
    sheet = normalize(path).sheet("Results")
    assert len(sheet.records) == 2

    record = sheet.records[0]
    assert len(record.assessments) == 4
    pairs = {
        (assessment.subject.normalized if assessment.subject else None,
         assessment.name.normalized,
         assessment.mark.normalized)
        for assessment in record.assessments
    }
    assert pairs == {
        ("physics", "ia", 10.0),
        ("physics", "see", 50.0),
        ("maths", "ia", 12.0),
        ("maths", "see", 40.0),
    }

    physics_ia = record.assessments[0]
    assert physics_ia.subject.raw == "Physics"
    assert physics_ia.name.raw == "IA"
    assert physics_ia.name.location.header_path == ["Physics", "IA"]
    assert physics_ia.name.location.column == 2
    assert physics_ia.mark.location.row == 3
    assert physics_ia.mark.location.render() == "Results!C3"


# ------------------------------------------------------------
# 4. Explicit total -> total, not subject_mark
# ------------------------------------------------------------


def test_explicit_total_is_not_subject_mark(tmp_path):
    path = write_workbook(
        tmp_path / "total.xlsx",
        {
            "Results": [
                ["Subject", "Total Marks", "Percentage"],
                ["Physics", 80, 80.0],
                ["Mathematics", 85, 85.0],
            ]
        },
    )
    sheet = normalize(path).sheet("Results")
    record = sheet.records[0]
    assert record.subject is not None
    assert record.subject_mark is None
    assert record.assessments == []

    total = record.find("total_marks")
    assert total.value.kind == "total"
    assert total.value.normalized == 80.0
    assert record.find("percentage").value.normalized == 80.0


def test_fraction_total_form(tmp_path):
    path = write_workbook(
        tmp_path / "total_fraction.xlsx",
        {
            "Results": [
                ["Subject", "Total Marks", "Percentage"],
                ["Physics", "80/100", 0.8],
            ]
        },
    )
    sheet = normalize(path).sheet("Results")
    record = sheet.records[0]
    total = record.find("total_marks")
    assert total.value.normalized == 80.0
    assert total.value.extra == {"maximum": 100.0}
    assert record.find("percentage").value.normalized == 80.0


# ------------------------------------------------------------
# 5. Status tokens are preserved, never zeroed
# ------------------------------------------------------------


def test_status_values_preserved_not_zeroed(tmp_path):
    path = write_workbook(
        tmp_path / "status.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Grade"],
                ["1", "Physics", "AB", "D"],
                ["2", "Maths", "Absent", "F"],
                ["3", "Chemistry", "Exempt", "B"],
                ["4", "English", 0, "B"],
                ["5", "Computer Science", 18, "B"],
                ["6", "Practical", "", "D"],
            ]
        },
    )
    sheet = normalize(path).sheet("Results")
    records = sheet.records
    assert len(records) == 6

    def _mark(record):
        return record.subject_mark.value

    assert _mark(records[0]).kind == VALUE_STATUS
    assert _mark(records[0]).normalized == "absent"
    assert _mark(records[0]).raw == "AB"

    assert _mark(records[1]).kind == VALUE_STATUS
    assert _mark(records[1]).normalized == "absent"

    assert _mark(records[2]).kind == VALUE_STATUS
    assert _mark(records[2]).normalized == "exempt"
    assert _mark(records[2]).raw == "Exempt"

    assert _mark(records[3]).kind == VALUE_ZERO
    assert _mark(records[3]).normalized == 0.0
    assert _mark(records[3]).raw == 0

    assert _mark(records[4]).kind == VALUE_MARK
    assert _mark(records[4]).normalized == 18.0

    assert _mark(records[5]).kind == VALUE_BLANK


# ------------------------------------------------------------
# 6. Percentage forms
# ------------------------------------------------------------


def test_percentage_forms_normalized(tmp_path):
    path = write_workbook(
        tmp_path / "percentages.xlsx",
        {"Data": [["Percentage"], ["69%"], [0.69], [69]]},
    )
    sheet = normalize(path).sheet("Data")
    records = sheet.records
    assert len(records) == 3

    percentages = [record.find("percentage").value for record in sheet.records]
    assert [value.normalized for value in percentages] == [69.0, 69.0, 69.0]
    assert [value.raw for value in percentages] == ["69%", 0.69, 69]
    # distinguishable: the fraction form is flagged, the "%" keeps its suffix
    assert "fraction" in percentages[1].note
    assert "%" in percentages[0].note


# ------------------------------------------------------------
# 7. Provenance
# ------------------------------------------------------------


def test_provenance_retained(tmp_path):
    path = write_workbook(
        tmp_path / "provenance.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                ["1", "Physics", 76, 76.0, "A"],
            ]
        },
    )
    normalization = normalize(path)
    assert normalization.source == str(path)
    assert normalization.filename == "provenance.xlsx"

    record = normalization.records[0]
    mark = record.subject_mark.value
    assert mark.location is not None
    assert mark.location.sheet == "Results"
    assert mark.location.row == 2
    assert mark.location.column == 2
    assert mark.location.header_text == "Marks"
    assert mark.location.render() == "Results!C2"

    grade = record.find("grade").value
    assert grade.location.render() == "Results!E2"

    subject = record.subject.value
    assert subject.location.render() == "Results!B2"


def test_provenance_preserves_multi_row_header_path(tmp_path):
    path = write_workbook(
        tmp_path / "provenance_multi.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Physics", "Physics", "Maths", "Maths"],
                ["", "", "IA", "SEE", "IA", "SEE"],
                ["1", "Aarav", 10, 50, 12, 40],
            ]
        },
    )
    normalization = normalize(path)
    assert len(normalization.records) == 1
    assessment = normalization.records[0].assessments[0]
    assert assessment.name.location.header_path == ["Physics", "IA"]
    assert assessment.name.location.column == 2
    assert assessment.name.location.header_text == "IA"
    assert assessment.mark.location.render() == "Results!C3"


# ------------------------------------------------------------
# 8. Ambiguity is preserved, never guessed
# ------------------------------------------------------------


def test_ambiguous_measure_preserved_not_guessed(tmp_path):
    path = write_workbook(
        tmp_path / "ambiguous.xlsx",
        {"Data": [["Marks"], [75], [80]]},
    )
    normalization = normalize(path)
    assert normalization.unresolved
    unresolved = normalization.unresolved[0]
    assert unresolved.code == "ambiguous_measure"
    assert len(unresolved.options) >= 2
    assert unresolved.sheet == "Data"

    sheet = normalization.sheet("Data")
    assert len(sheet.records) == 2
    for record in sheet.records:
        # the normalizer must not guess a meaning for the marks column
        assert record.unresolved
        value = record.unresolved[0].value
        assert value.kind == VALUE_UNRESOLVED
        assert value.normalized is None
        assert value.location.render() == f"Data!A{record.row}"
    assert sheet.records[0].unresolved[0].value.raw == 75
    assert sheet.records[1].unresolved[0].value.raw == 80


# ------------------------------------------------------------
# 9. Empty and sparse sheets
# ------------------------------------------------------------


def test_empty_sheet(tmp_path):
    workbook = Workbook()
    workbook.active.title = "Blank"
    path = tmp_path / "empty.xlsx"
    workbook.save(path)

    sheet = normalize(path).sheet("Blank")
    assert sheet.empty is True
    assert sheet.records == []
    assert sheet.unresolved == []


def test_sparse_sheet_with_header_only(tmp_path):
    path = write_workbook(
        tmp_path / "sparse.xlsx",
        {"Results": [["S.No.", "Subject", "Marks", "Grade"]]},
    )
    sheet = normalize(path).sheet("Results")
    assert sheet.empty is False
    assert sheet.records == []
    assert any("sparse_sheet" in warning for warning in sheet.warnings)


# ------------------------------------------------------------
# 10. Real semester-result-like structure (no hardcoded names)
# ------------------------------------------------------------


def test_semester_result_like_workbook(tmp_path):
    rows = [
        ["FIRST SEMESTER ECE RESULT", "", "", "", ""],
        ["", "", "", "", ""],
        ["Student Name", "Quaja Hussain", "", "", ""],
        ["Semester", "First Sem", "", "", ""],
        ["Total Marks", "69/100", "", "", ""],
        ["Overall Percentage", "69.00%", "", "", ""],
        ["", "", "", "", ""],
        ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
        [1, "AI", 69, 0.69, "B"],
    ]
    path = write_workbook(tmp_path / "term_report.xlsx", {"Term Report": rows})

    sheet = normalize(path).sheet("Term Report")
    assert sheet.header_rows == [8]
    assert len(sheet.records) == 1

    record = sheet.records[0]
    assert record.subject.value.normalized == "AI"
    assert record.subject_mark.value.kind == VALUE_MARK
    assert record.subject_mark.value.normalized == 69.0
    assert record.subject_mark.value.raw == 69
    assert record.find("percentage").value.normalized == 69.0
    assert record.find("percentage").value.raw == 0.69
    assert record.find("grade").value.normalized == "B"
    assert record.assessments == []
    assert sheet.unresolved == []


def test_three_semester_result_like_workbooks(tmp_path):
    structures = [
        {
            "name": "Quaja",
            "rows": [
                ["Student Name", "Quaja Hussain", "", "", ""],
                ["Total Marks", "69/100", "", "", ""],
                ["Overall Percentage", "69.00%", "", "", ""],
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                [1, "AI", 69, 0.69, "B"],
            ],
        },
        {
            "name": "Priya",
            "rows": [
                ["Student Name", "Priya Sharma", "", "", ""],
                ["Total Marks", "404/500", "", "", ""],
                ["Overall Percentage", "80.80%", "", "", ""],
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                [1, "Mathematics", 94, 0.94, "A+"],
                [2, "Physics", 82, 0.82, "A"],
                [3, "Chemistry", 68, 0.68, "B"],
            ],
        },
        {
            "name": "Mohammed",
            "rows": [
                ["Student Name", "M.S. Mohammed Quaja Hussain", "", "", ""],
                ["Total Marks", "410/500", "", "", ""],
                ["Overall Percentage", "82.00%", "", "", ""],
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                [1, "Maths", 80, 0.8, "A"],
                [2, "Python", 80, 0.8, "A"],
                [3, "DT", 90, 0.9, "A+"],
            ],
        },
    ]
    for structure in structures:
        path = write_workbook(
            tmp_path / f"{structure['name'].lower()}_result.xlsx",
            {"Result": structure["rows"]},
        )
        sheet = normalize(path).sheet("Result")
        assert sheet.records, structure["name"]
        for record in sheet.records:
            assert record.subject is not None, structure["name"]
            assert record.subject_mark is not None, structure["name"]
            assert record.assessments == [], structure["name"]
            percentage = record.find("percentage").value
            assert percentage.normalized == pytest.approx(
                percentage.raw * 100, abs=0.01
            ), structure["name"]


# ------------------------------------------------------------
# 11. Determinism and public API
# ------------------------------------------------------------


def test_normalization_is_deterministic(tmp_path):
    path = write_workbook(
        tmp_path / "deterministic.xlsx",
        {
            "Data": [
                ["Roll No", "Student Name", "Physics IA", "Maths SEE", "Total"],
                ["1", "Aarav", 20, 30, 50],
                ["2", "Bhavna", 18, 28, 46],
            ]
        },
    )
    first = normalize(path)
    second = normalize(path)
    assert first == second


def test_public_api_shape(tmp_path):
    path = write_workbook(
        tmp_path / "api.xlsx",
        {"Data": [["Subject", "Marks", "Grade"], ["Physics", 45, "C"]]},
    )
    profile = inspect_workbook(path)
    detection = detect_workbook(profile)
    normalization = normalize_workbook_profile(profile, detection=detection)
    assert isinstance(normalization, WorkbookNormalization)
    assert normalization.filename == "api.xlsx"
    assert normalization.records
    assert isinstance(normalization.records[0].subject_mark.value, NormalizedValue)

    # normalize_sheet works on an inspected worksheet too
    from openpyxl import Workbook as Wb

    wb = Wb()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Subject", "Marks"])
    ws.append(["Physics", 45])
    sheet_profile = inspect_sheet(ws)
    sheet_semantics = detect_workbook(inspect_workbook(path)).sheets[0]
    rows = {1: ("Subject", "Marks"), 2: ("Physics", 45)}
    normalized_sheet = normalize_sheet(sheet_profile, sheet_semantics, rows)
    assert len(normalized_sheet.records) == 1