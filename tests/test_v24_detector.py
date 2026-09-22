"""Tests for V2.4 Phase 2: the Intelligent Detector.

The detector interprets a Phase 1 structural profile into semantic
candidates - identity, subjects, assessments and subject-assessment
pairs - carrying evidence and a deterministic, bounded confidence for
every reading.  It must never silently resolve ambiguity: columns whose
meaning cannot be decided from the evidence produce an Ambiguity record
instead of a confident guess, and conflicting evidence is preserved as a
Conflict.  These tests verify the semantic readings (not the structure,
which Phase 1 already covers) plus the determinism and explainability
guarantees.
"""

from assessment_engine.config import load_config
import pytest
from assessment_engine.intelligent_import import (
    Ambiguity,
    Conflict,
    FAMILY_ACADEMIC,
    FAMILY_ASSESSMENT,
    FAMILY_IDENTITY,
    FAMILY_ORGANIZATIONAL,
    FAMILY_SUBJECT,
    FAMILY_SUBJECT_ASSESSMENT,
    SemanticCandidate,
    SheetSemantics,
    Vocabulary,
    detect_sheet,
    detect_workbook,
    inspect_workbook,
    vocabulary_from_config,
)

config_path = "configs/college.json"


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def write_workbook(path, sheets):
    """sheets: dict {sheet_name: list_of_rows}."""
    from openpyxl import Workbook

    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


def detect(path):
    return detect_workbook(inspect_workbook(path))


# ------------------------------------------------------------
# Public API
# ------------------------------------------------------------


def test_detector_public_api(tmp_path):
    path = write_workbook(
        tmp_path / "api.xlsx",
        {"Data": [["Roll No", "Physics", "IA"], ["1", 20, 10]]},
    )
    profile = inspect_workbook(path)

    assert detect_workbook(profile).filename == "api.xlsx"
    assert detect_sheet(profile.sheet("Data")).name == "Data"
    assert Vocabulary() is not None
    assert isinstance(vocabulary_from_config(load_config(config_path)), Vocabulary)


# ------------------------------------------------------------
# Identity
# ------------------------------------------------------------


@pytest.mark.parametrize("label", ["Roll No", "Roll Number", "Student Name", "Student ID", "Registration Number"])
def test_identity_labels(tmp_path, label):
    path = write_workbook(
        tmp_path / "identity.xlsx",
        {"Data": [[label, "Physics"], ["2024-001", 45]]},
    )
    sheet = detect(path).sheets[0]
    assert sheet.identity, f"no identity interpretation for {label!r}"
    assert sheet.identity[0].family == FAMILY_IDENTITY
    assert sheet.identity[0].confidence > 0.5
    assert all(e.source and e.explanation for e in sheet.identity[0].evidence)


# ------------------------------------------------------------
# Assessment terms
# ------------------------------------------------------------


@pytest.mark.parametrize("label", ["IA", "SEE", "Midterm", "Internal", "External", "Practical", "Assignment"])
def test_assessment_term_columns(tmp_path, label):
    path = write_workbook(
        tmp_path / "assessment.xlsx",
        {"Data": [[label], [20], [18]]},
    )
    sheet = detect(path).sheets[0]
    column = sheet.columns[0]
    assert column.best is not None
    assert column.best.family in (FAMILY_ASSESSMENT, FAMILY_ACADEMIC)
    assert any(c.family == FAMILY_ASSESSMENT for c in column.candidates)


# ------------------------------------------------------------
# Subjects vs non-subjects on a subject-looking workflow
# ------------------------------------------------------------


def test_obvious_subject_names(tmp_path):
    path = write_workbook(
        tmp_path / "subjects.xlsx",
        {
            "Data": [
                ["Roll No", "Physics", "Mathematics", "Chemistry"],
                ["1", 45, 48, 40],
                ["2", 50, 52, 44],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    subject_columns = [
        c.header_text for c in sheet.columns if c.best and c.best.family == FAMILY_SUBJECT
    ]
    assert subject_columns == ["Physics", "Mathematics", "Chemistry"]


def test_obvious_non_subject_columns_never_subjects(tmp_path):
    path = write_workbook(
        tmp_path / "non_subjects.xlsx",
        {
            "Data": [
                [
                    "Roll No", "Student Name", "Internal", "External", "Final",
                    "Semester", "Year", "Total", "Percentage", "Average", "Rank",
                    "Attendance", "Grade", "Result", "Status",
                ],
                ["1", "Aarav", 20, 60, 40, 1, 2026, 80, 80.0, 40.0, 1, 90, "A", "P", "Pass"],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    non_subject = [c.header_text for c in sheet.columns if c.best and c.best.family == FAMILY_SUBJECT]
    assert non_subject == []
    assert sheet.subjects == []


# ------------------------------------------------------------
# Combined subject-assessment headers
# ------------------------------------------------------------


def test_combined_subject_assessment_headers(tmp_path):
    path = write_workbook(
        tmp_path / "combined.xlsx",
        {
            "Results": [
                ["Roll No", "Physics IA", "Physics SEE", "Maths IA", "Maths SEE"],
                ["6A01", 78, 65, 82, 70],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    pairs = {(p.subject, p.assessment) for p in sheet.pairs}
    assert pairs == {
        ("physics", "ia"),
        ("physics", "see"),
        ("maths", "ia"),
        ("maths", "see"),
    }
    for column in sheet.columns[1:]:
        assert column.best.family == FAMILY_SUBJECT_ASSESSMENT
        assert column.best.subject and column.best.assessment
        assert column.best.confidence > 0.9


# ------------------------------------------------------------
# Multi-row headers (hierarchical interpretation)
# ------------------------------------------------------------


def test_multi_row_header_hierarchy(tmp_path):
    path = write_workbook(
        tmp_path / "multi.xlsx",
        {
            "Data": [
                ["Roll No", "Student Name", "Physics", "Physics", "Maths", "Maths"],
                ["", "", "IA", "SEE", "IA", "SEE"],
                ["1", "Aarav", 10, 50, 12, 40],
                ["2", "Bhavna", 11, 52, 13, 41],
            ]
        },
    )
    profile = inspect_workbook(path)
    assert profile.sheet("Data").header_block.is_multi_row

    sheet = detect_workbook(profile).sheets[0]
    assert sheet.columns[2].header_path == ["Physics", "IA"]
    assert sheet.columns[3].header_path == ["Physics", "SEE"]
    assert sheet.columns[4].header_path == ["Maths", "IA"]
    assert sheet.columns[5].header_path == ["Maths", "SEE"]

    pairs = {(p.subject, p.assessment) for p in sheet.pairs}
    assert pairs == {
        ("physics", "ia"),
        ("physics", "see"),
        ("maths", "ia"),
        ("maths", "see"),
    }
    for column in sheet.columns[2:]:
        assert column.best.family == FAMILY_SUBJECT_ASSESSMENT
        assert any(
            e.signal == "header_hierarchy" for e in column.best.evidence
        )


# ------------------------------------------------------------
# Sheet-name evidence
# ------------------------------------------------------------


def test_sheet_name_subject_evidence(tmp_path):
    path = write_workbook(
        tmp_path / "sheetname.xlsx",
        {"Physics": [["Roll No", "IA"], ["1", 20], ["2", 18]]},
    )
    sheet = detect(path).sheets[0]
    subjects = [c for c in sheet.sheet_subjects]
    assert len(subjects) == 1
    assert subjects[0].normalized == "physics"
    assert any(e.source == "sheet_name" for e in subjects[0].evidence)


def test_sheet_name_with_program_stripped(tmp_path):
    path = write_workbook(
        tmp_path / "prog.xlsx",
        {"DSA Program": [["Roll No", "IA"], ["1", 20]]},
    )
    sheet = detect(path).sheets[0]
    assert [c.normalized for c in sheet.sheet_subjects] == ["dsa"]


def test_cross_sheet_consistency(tmp_path):
    path = write_workbook(
        tmp_path / "cross.xlsx",
        {
            "Physics": [["Roll No", "IA", "Total"], ["1", 20, 80]],
            "Physics Results": [["Roll No", "SEE", "Total"], ["2", 30, 70]],
        },
    )
    detection = detect(path)
    assert [c.normalized for c in detection.cross_sheet_subjects] == ["physics"]
    for candidate in detection.cross_sheet_subjects:
        assert any(e.source == "cross_sheet" for e in candidate.evidence)


# ------------------------------------------------------------
# Ambiguity is preserved, never silently resolved
# ------------------------------------------------------------


def test_generic_id_stays_ambiguous(tmp_path):
    path = write_workbook(
        tmp_path / "id.xlsx",
        {"Data": [["ID", "Physics"], ["1", 45]]},
    )
    sheet = detect(path).sheets[0]
    id_column = sheet.columns[0]
    assert id_column.best.confidence < 0.65
    codes = [a.code for a in id_column.ambiguity]
    assert "id_column" in codes
    id_ambiguity = next(a for a in id_column.ambiguity if a.code == "id_column")
    assert len(id_ambiguity.options) >= 2


def test_final_column_is_ambiguous(tmp_path):
    path = write_workbook(
        tmp_path / "final.xlsx",
        {"Data": [["Final", "Total"], ["A", 80], ["B", 75]]},
    )
    sheet = detect(path).sheets[0]
    final_column = sheet.columns[0]
    assert "final_column" in [a.code for a in final_column.ambiguity]


def test_marks_column_is_ambiguous(tmp_path):
    path = write_workbook(
        tmp_path / "marks.xlsx",
        {"Data": [["Marks"], [75], [80]]},
    )
    sheet = detect(path).sheets[0]
    marks_column = sheet.columns[0]
    assert "ambiguous_measure" in [a.code for a in marks_column.ambiguity]


# ------------------------------------------------------------
# Contextual column relationships
#
# A generic "Marks" column is resolved by reasoning about the columns
# around it, not by classifying the header in isolation.
# ------------------------------------------------------------


def test_subject_result_table_marks_is_subject_level(tmp_path):
    path = write_workbook(
        tmp_path / "subject_result.xlsx",
        {
            "Results": [
                ["S.No.", "Subject", "Marks", "Percentage", "Grade"],
                ["1", "Physics", 76, 76.0, "A"],
                ["2", "Mathematics", 82, 82.0, "A+"],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    by_header = {c.header_text: c for c in sheet.columns}

    subject_col = by_header["Subject"]
    assert subject_col.best.family == FAMILY_SUBJECT
    assert subject_col.best.kind == "subject"

    marks_col = by_header["Marks"]
    assert marks_col.best is not None
    assert marks_col.best.family == FAMILY_SUBJECT
    assert marks_col.best.kind == "subject_mark"
    assert marks_col.best.confidence > 0.7
    assert marks_col.best.subject == "Subject"

    assessment = next(
        (c for c in marks_col.candidates if c.kind == "marks_obtained"), None
    )
    assert assessment is not None
    assert marks_col.best.confidence > assessment.confidence

    assert "ambiguous_measure" not in [a.code for a in marks_col.ambiguity]
    assert any(
        e.signal == "subject_column_cooccurrence" for e in marks_col.best.evidence
    )
    assert any(
        e.signal == "assessment_column_relation" for e in assessment.evidence
    )

    grade_col = by_header["Grade"]
    assert grade_col.best is not None and grade_col.best.kind == "grade"
    percentage_col = by_header["Percentage"]
    assert percentage_col.best is not None and percentage_col.best.kind == "percentage"


def test_assessment_table_columns_keep_assessment_meaning(tmp_path):
    path = write_workbook(
        tmp_path / "assessment_table.xlsx",
        {
            "Results": [
                ["Roll No", "Student Name", "Internal", "External"],
                ["1", "Aarav", 20, 60],
                ["2", "Bhavna", 22, 58],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    by_header = {c.header_text: c for c in sheet.columns}

    for header in ("Internal", "External"):
        column = by_header[header]
        assert column.best.family == FAMILY_ASSESSMENT
        assert any(c.family == FAMILY_ASSESSMENT for c in column.candidates)

    for column in sheet.columns:
        assert not any(c.kind == "subject_mark" for c in column.candidates)

    assert sheet.identity[0].kind == "roll_no"


def test_total_marks_stays_total_not_subject(tmp_path):
    path = write_workbook(
        tmp_path / "total_marks.xlsx",
        {
            "Results": [
                ["Subject", "Total Marks", "Percentage"],
                ["Physics", 80, 80.0],
                ["Mathematics", 85, 85.0],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    by_header = {c.header_text: c for c in sheet.columns}

    total_col = by_header["Total Marks"]
    assert total_col.best is not None
    assert total_col.best.kind == "total_marks"
    assert not any(c.kind == "subject_mark" for c in total_col.candidates)


def test_final_marks_keeps_final_semantics(tmp_path):
    path = write_workbook(
        tmp_path / "final_marks.xlsx",
        {
            "Results": [
                ["Subject", "Final Marks", "Grade"],
                ["Physics", 78, "A"],
                ["Mathematics", 82, "A+"],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    by_header = {c.header_text: c for c in sheet.columns}

    final_col = by_header["Final Marks"]
    assert final_col.best is not None
    assert final_col.best.family == FAMILY_ASSESSMENT
    assert all(c.kind != "subject_mark" for c in final_col.candidates)


def test_subject_assessment_table_relationship_preserved(tmp_path):
    path = write_workbook(
        tmp_path / "subject_assessment.xlsx",
        {
            "Results": [
                ["Subject", "Internal", "External"],
                ["Physics", 20, 60],
                ["Mathematics", 22, 58],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    by_header = {c.header_text: c for c in sheet.columns}

    assert by_header["Subject"].best.kind == "subject"
    for header in ("Internal", "External"):
        column = by_header[header]
        assert column.best.family == FAMILY_ASSESSMENT
        assert any(c.family == FAMILY_ASSESSMENT for c in column.candidates)
        assert all(c.kind != "subject_mark" for c in column.candidates)

    assert sheet.subjects, "the Subject column keeps the subject reading"
    assert not any(note.startswith("assessment_without_subject") for note in sheet.notes)


def test_ambiguity_records_carry_sheet(tmp_path):
    path = write_workbook(
        tmp_path / "amb_sheet.xlsx",
        {"Results": [["ID", "Final"], ["1", "A"]]},
    )
    detection = detect(path)
    assert detection.ambiguity
    for entry in detection.ambiguity:
        assert isinstance(entry, Ambiguity)
        assert entry.sheet == "Results"


# ------------------------------------------------------------
# Conflicts are preserved
# ------------------------------------------------------------


def test_name_header_with_numeric_data_conflicts(tmp_path):
    path = write_workbook(
        tmp_path / "conflict_name.xlsx",
        {"Data": [["Student Name", "Total"], [1, 80], [2, 75]]},
    )
    sheet = detect(path).sheets[0]
    name_column = sheet.columns[0]
    assert len(name_column.conflicts) == 1
    assert name_column.conflicts[0].code == "header_data_conflict"


def test_subject_above_measure_header_conflicts(tmp_path):
    path = write_workbook(
        tmp_path / "conflict_measure.xlsx",
        {
            "Data": [
                ["Roll No", "Physics", "Maths"],
                ["", "Total", "Average"],
                ["1", 80, 75],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    measure_columns = sheet.columns[1:]
    assert [c.conflicts[0].code for c in measure_columns] == [
        "conflicting_measure",
        "conflicting_measure",
    ]
    for column in measure_columns:
        assert isinstance(column.conflicts[0], Conflict)


def test_unambiguous_sheet_has_no_conflicts(tmp_path):
    path = write_workbook(
        tmp_path / "clean.xlsx",
        {"Data": [["Roll No", "Physics", "Total"], ["1", 45, 45]]},
    )
    detection = detect(path)
    assert detection.conflicts == []
    assert detection.ambiguity == []


# ------------------------------------------------------------
# Confidence: deterministic, bounded, evidence-backed
# ------------------------------------------------------------


def test_confidence_is_bounded_and_evidence_backed(tmp_path):
    path = write_workbook(
        tmp_path / "bounds.xlsx",
        {
            "Data": [
                ["Roll No", "Physics IA", "Physics SEE", "Total"],
                ["1", 45, 30, 75],
            ]
        },
    )
    sheet = detect(path).sheets[0]
    for candidate in [c.best for c in sheet.columns if c.best]:
        assert 0.0 <= candidate.confidence <= 0.99
        assert candidate.evidence
        assert candidate.score == round(sum(e.weight for e in candidate.evidence), 3)


def test_detection_is_deterministic(tmp_path):
    path = write_workbook(
        tmp_path / "deterministic.xlsx",
        {
            "Data": [
                ["Roll No", "Student Name", "Physics IA", "Maths SEE", "Total"],
                ["1", "Aarav", 20, 30, 50],
                ["2", "Bhavna", 18, 28, 46],
            ]
        },
    )
    profile = inspect_workbook(path)
    first = detect_workbook(profile)
    second = detect_workbook(profile)
    first_models = [
        (c.best.family, c.best.kind, c.best.confidence, c.best.score)
        for sheet in first.sheets
        for c in sheet.columns
        if c.best
    ]
    second_models = [
        (c.best.family, c.best.kind, c.best.confidence, c.best.score)
        for sheet in second.sheets
        for c in sheet.columns
        if c.best
    ]
    assert first_models == second_models


def test_more_evidence_raises_strength(tmp_path):
    from assessment_engine.intelligent_import.detector.evidence import Evidence

    single = [Evidence("h", "t", "ia", 0.90, "one")]
    double = single + [Evidence("d", "d", "ia", 0.20, "data")]
    from assessment_engine.intelligent_import.detector.evidence import candidate_confidence

    assert candidate_confidence(double) > candidate_confidence(single)
    assert candidate_confidence([Evidence("h", "t", "x", 0.90, "one")]) < 1.0


# ------------------------------------------------------------
# Institution vocabulary extension point
# ------------------------------------------------------------


def test_vocabulary_from_config_extracts_terms():
    vocab = vocabulary_from_config(load_config(config_path))
    assert "internal" in vocab.assessment_terms
    assert "external" in vocab.assessment_terms
    assert "data structures and algorithms" in vocab.subject_terms
    assert "dbms" in vocab.subject_terms


def test_detector_uses_config_backed_vocabulary(tmp_path):
    path = write_workbook(
        tmp_path / "vocab.xlsx",
        {
            "March": [
                ["Roll No", "Data Structures and Algorithms", "DBMS", "Internal"],
                ["1", 45, 42, 20],
                ["2", 50, 40, 18],
            ]
        },
    )
    cfg = load_config(config_path)
    sheet = detect_workbook(inspect_workbook(path), config=cfg).sheets[0]
    for column in sheet.columns[1:3]:
        candidate = column.best
        assert candidate is not None and candidate.family == FAMILY_SUBJECT
        assert any(
            e.source == "institution_vocabulary" for e in candidate.evidence
        )


def test_detector_accepts_explicit_vocabulary(tmp_path):
    path = write_workbook(
        tmp_path / "vocab_explicit.xlsx",
        {"Data": [["Roll No", "Biophysics"], ["1", 40]]},
    )
    vocab = Vocabulary(
        subject_terms=["biophysics"],
        subject_aliases={"biophysics": "Biophysics"},
    )
    sheet = detect_workbook(inspect_workbook(path), vocabulary=vocab).sheets[0]
    subject_column = sheet.columns[1]
    assert subject_column.best.family == FAMILY_SUBJECT
    assert any(
        e.source == "institution_vocabulary" for e in subject_column.best.evidence
    )


# ------------------------------------------------------------
# Semantics shape and empty sheets
# ------------------------------------------------------------


def test_empty_sheet_produces_empty_semantics(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Blank"
    path = tmp_path / "blank.xlsx"
    workbook.save(path)

    sheet = detect(path).sheets[0]
    assert isinstance(sheet, SheetSemantics)
    assert sheet.empty is True
    assert sheet.columns == []
    assert sheet.identity == []
    assert sheet.assessments == []
    assert sheet.ambiguity == []


def test_assessment_without_subject_note(tmp_path):
    path = write_workbook(
        tmp_path / "no_subject.xlsx",
        {"Results": [["Roll No", "IA", "SEE"], ["1", 20, 60]]},
    )
    sheet = detect(path).sheets[0]
    assert any(note.startswith("assessment_without_subject") for note in sheet.notes)


def test_workbook_aggregates_sheet_records(tmp_path):
    path = write_workbook(
        tmp_path / "aggregate.xlsx",
        {
            "A": [["ID", "Physics"], ["1", 45]],
            "B": [["Final", "Maths"], ["A", 48]],
        },
    )
    detection = detect(path)
    assert len(detection.sheets) == 2
    assert any(
        a.code == "id_column" for a in detection.ambiguity
    )
    assert any(
        a.code == "final_column" for a in detection.ambiguity
    )


# ------------------------------------------------------------
# Matching primitives shared with later phases
# ------------------------------------------------------------


def test_identity_matches_primitive():
    from assessment_engine.intelligent_import.detector import identity_matches

    specs = identity_matches("roll no")
    assert specs and specs[0].family == FAMILY_IDENTITY
    assert specs[0].kind == "roll_no"


def test_assessments_in_splits_subject_remainder():
    from assessment_engine.intelligent_import.detector import assessments_in

    split = assessments_in("physics ia")
    assert split.terms == [("ia", 0.90)]
    assert split.remainder == ["physics"]

    split = assessments_in("internal assessment term 1")
    terms = [t for t, _ in split.terms]
    assert "internal assessment" in terms
    assert all(not token.isdigit() for token in split.remainder)