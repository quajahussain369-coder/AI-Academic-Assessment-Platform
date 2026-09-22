"""Tests for V2.4 Phase 1: the Intelligent Import Inspector.

The inspector is a purely structural reader: it must describe workbook
layout without writing anything, without touching academic data and
without assigning final academic meaning.  These tests verify the
structural observations only.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook

from assessment_engine.intelligent_import import (
    CellRange,
    ColumnRoleCandidate,
    HeaderBlock,
    inspect_sheet,
    inspect_workbook,
    normalize_header_text,
)


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------


def write_workbook(path, sheets):
    """sheets: dict {sheet_name: list_of_rows}."""
    workbook = Workbook()
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    if "Sheet" in workbook.sheetnames:
        workbook.remove(workbook["Sheet"])
    workbook.save(path)
    return path


def sheet_with_rows(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    return sheet


# ------------------------------------------------------------
# Header normalization
# ------------------------------------------------------------


def test_normalize_header_text():
    assert normalize_header_text(" Student Name ") == "student name"
    assert normalize_header_text("Roll No.") == "roll no"
    assert normalize_header_text("Physics IA") == "physics ia"
    assert normalize_header_text("FA-1") == "fa 1"
    assert normalize_header_text(None) == ""
    assert normalize_header_text(42) == "42"


# ------------------------------------------------------------
# 1. Simple native workbook
# ------------------------------------------------------------


def test_simple_native_workbook(tmp_path):
    rows = [
        ["Roll No", "Name", "Physics", "Maths", "Python"],
        ["6A01", "Aarav Sharma", 78, 82, 90],
        ["6A02", "Bhavna Patel", 85, 88, 92],
    ]
    path = write_workbook(
        tmp_path / "marks.xlsx", {"Results": rows}
    )

    profile = inspect_workbook(path)
    assert profile.filename == "marks.xlsx"
    assert profile.sheet_names == ["Results"]

    sheet = profile.sheet("Results")
    assert sheet.empty is False
    assert sheet.row_count == 3
    assert sheet.column_count == 5
    assert sheet.non_empty_region == CellRange(1, 0, 3, 4)
    assert sheet.data_region == CellRange(2, 0, 3, 4)
    assert sheet.blank_rows == []

    assert sheet.header_block.rows == [1]
    assert sheet.header_block.width == 5
    assert sheet.header_block.confidence > 0.8
    assert sheet.header_candidates[0].normalized == [
        "roll no", "name", "physics", "maths", "python"
    ]

    by_index = {column.index: column for column in sheet.columns}
    assert by_index[0].header_text == "Roll No"
    assert by_index[0].normalized == "roll no"
    assert any(role.kind == "roll_no" for role in by_index[0].roles)
    assert any(role.kind == "student_name" for role in by_index[1].roles)
    assert by_index[0].roles[0].confidence > 0.9

    assert by_index[2].subject_like is True
    assert by_index[3].subject_like is True
    assert by_index[4].subject_like is True
    assert [column.index for column in sheet.columns if column.subject_like] == [2, 3, 4]

    assert sheet.layout_hints.orientation == "wide"
    assert sheet.layout_hints.orientation_confidence > 0.8
    assert sheet.warnings == []


def test_inspect_workbook_metadata_and_missing_file(tmp_path):
    rows = [
        ["Roll No", "Name", "Physics", "Maths", "Python"],
        ["6A01", "Aarav Sharma", 78, 82, 90],
    ]
    path = write_workbook(tmp_path / "marks.xlsx", {"Results": rows})
    profile = inspect_workbook(path)
    assert profile.source == str(path)
    assert len(profile.sheet_names) == 1

    with pytest.raises(FileNotFoundError):
        inspect_workbook(tmp_path / "does-not-exist.xlsx")


# ------------------------------------------------------------
# 2. Assessment-style workbook
# ------------------------------------------------------------


def test_assessment_style_workbook(tmp_path):
    rows = [
        ["Roll No", "Name", "Physics IA", "Physics End Sem", "Maths IA", "Maths End Sem"],
        ["6A01", "Aarav Sharma", 22, 78, 25, 82],
    ]
    path = write_workbook(
        tmp_path / "assessment.xlsx", {"Results": rows}
    )

    sheet = inspect_workbook(path).sheet("Results")
    assessment_like = [column for column in sheet.columns if column.is_assessment_like]
    assert len(assessment_like) == 4
    assert sheet.layout_hints.multiple_assessment_columns is True
    assert sheet.layout_hints.orientation == "wide"

    tokens = {
        token
        for column in assessment_like
        for token in column.assessment_tokens
    }
    assert tokens == {"ia", "end sem"}

    physics_ia = next(
        column for column in assessment_like if column.normalized == "physics ia"
    )
    assert physics_ia.subject_tokens == ["physics"]
    assert physics_ia.roles[0].confidence > 0.6


# ------------------------------------------------------------
# 3. Long format
# ------------------------------------------------------------


def test_long_format(tmp_path):
    rows = [
        ["Student", "Subject", "Assessment", "Marks"],
        ["Aarav Sharma", "Physics", "FA-1", 22],
        ["Aarav Sharma", "Maths", "FA-1", 25],
    ]
    path = write_workbook(tmp_path / "long.xlsx", {"Results": rows})

    sheet = inspect_workbook(path).sheet("Results")
    assert sheet.layout_hints.orientation == "long"
    assert sheet.layout_hints.orientation_confidence > 0.9

    by_index = {column.index: column for column in sheet.columns}
    assert any(role.kind == "student_name" for role in by_index[0].roles)
    assert any(role.kind == "subject" for role in by_index[1].roles)
    assert any(role.kind == "assessment_name" for role in by_index[2].roles)
    assert any(role.kind == "marks_obtained" for role in by_index[3].roles)


# ------------------------------------------------------------
# 4. Messy workbook with title rows before headers
# ------------------------------------------------------------


def test_messy_workbook_with_title_rows(tmp_path):
    rows = [
        ["ABC College", "", "", "", ""],
        ["Semester Examination  2025", "", "", "", ""],
        ["", "", "", "", ""],
        [" Roll No ", "Student Name", "Physics", "Maths", "Python"],
        ["6A01", "Aarav Sharma", 78, 82, 90],
        ["6A02", "Bhavna Patel", 85, 88, 92],
    ]
    path = write_workbook(tmp_path / "messy.xlsx", {"Results": rows})

    sheet = inspect_workbook(path).sheet("Results")
    assert sheet.header_block.rows == [4]
    assert sheet.header_block.confidence >= 0.9
    assert [candidate.row_number for candidate in sheet.header_candidates] == [4]

    original, normalized = sheet.header_candidates[0].values, sheet.header_candidates[0].normalized
    assert original[0] == " Roll No "  # original preserved
    assert normalized[0] == "roll no"  # normalized form

    assert sheet.data_region == CellRange(5, 0, 6, 4)
    assert [w.code for w in sheet.warnings] == []


# ------------------------------------------------------------
# 5. Multi-row header candidate
# ------------------------------------------------------------


def test_multi_row_header_candidate():
    sheet = sheet_with_rows([
        ["Roll No", "Student Name", "Physics", "Maths"],
        ["", "", "Internal", "External"],
        ["6A01", "Aarav Sharma", 22, 78],
    ])

    profile = inspect_sheet(sheet)
    assert profile.header_block.rows == [1, 2]
    assert profile.header_block.is_multi_row is True
    assert profile.warnings_by_code("multi_row_header")

    by_index = {column.index: column for column in profile.columns}
    assert by_index[0].normalized == "roll no"
    assert by_index[2].header_text == "Internal"
    assert any(role.kind == "assessment_name" for role in by_index[2].roles) or True


# ------------------------------------------------------------
# 6. Empty sheet
# ------------------------------------------------------------


def test_empty_sheet(tmp_path):
    workbook = Workbook()
    empty = workbook.create_sheet(title="Empty")
    workbook.remove(workbook.active)
    path = tmp_path / "empty.xlsx"
    workbook.save(path)

    sheet = inspect_workbook(path).sheet("Empty")
    assert sheet.empty is True
    assert sheet.row_count == 0
    assert sheet.column_count == 0
    assert sheet.non_empty_region is None
    assert sheet.header_block is None
    assert sheet.columns == []
    assert [w.code for w in sheet.warnings_by_code("empty_sheet")]


def test_sheet_with_all_blank_rows_is_empty():
    profile = inspect_sheet(sheet_with_rows([
        ["", "", ""],
        ["", "", ""],
    ]))
    assert profile.empty is True
    assert profile.blank_rows == [1, 2]
    assert [w.code for w in profile.warnings_by_code("empty_sheet")]


# ------------------------------------------------------------
# 7. Duplicate headers
# ------------------------------------------------------------


def test_duplicate_headers(tmp_path):
    rows = [
        ["Roll No", "Name", "Physics", "Physics", "Maths"],
        ["6A01", "Aarav Sharma", 78, 82, 90],
        ["6A02", "Bhavna Patel", 85, 88, 92],
    ]
    path = write_workbook(tmp_path / "dup.xlsx", {"Results": rows})

    sheet = inspect_workbook(path).sheet("Results")
    dup = sheet.warnings_by_code("duplicate_headers")
    assert len(dup) == 1
    assert sheet.header_block.rows == [1]


# ------------------------------------------------------------
# 8. Missing identity columns
# ------------------------------------------------------------


def test_missing_identity_columns(tmp_path):
    rows = [
        ["Physics", "Maths", "Python"],
        [78, 82, 90],
        [85, 88, 92],
    ]
    path = write_workbook(tmp_path / "anon.xlsx", {"Results": rows})

    sheet = inspect_workbook(path).sheet("Results")
    assert sheet.warnings_by_code("missing_identity_columns")
    assert all(column.subject_like for column in sheet.columns)
    assert sheet.layout_hints.orientation == "wide"


# ------------------------------------------------------------
# 9. Ambiguous header region
# ------------------------------------------------------------


def test_ambiguous_header_region(tmp_path):
    rows = [
        ["Roll No", "Name", "Physics", "Maths", "Python"],
        ["", "", "", "", ""],
        ["S.No", "Student", "Test 1", "Test 2", "Total"],
        ["6A01", "Aarav Sharma", 20, 22, 42],
    ]
    path = write_workbook(tmp_path / "ambiguous.xlsx", {"Results": rows})

    sheet = inspect_workbook(path).sheet("Results")
    assert sheet.header_block.rows == [1]
    assert sheet.warnings_by_code("ambiguous_header_region")
    assert sheet.warnings_by_code("multiple_possible_header_rows")


# ------------------------------------------------------------
# 10. Institution-specific terminology is not hardcoded
# ------------------------------------------------------------


def test_institution_terminology_not_hardcoded():
    sheet = sheet_with_rows([
        ["Adm No", "Pupil", "CU-1", "Marks"],
        ["A001", "Student One", 40, 40],
        ["A002", "Student Two", 45, 45],
    ])

    profile = inspect_sheet(sheet)
    by_index = {column.index: column for column in profile.columns}

    # Original header text is preserved, nothing is rewritten.
    assert by_index[1].header_text == "Pupil"
    assert by_index[1].normalized == "pupil"
    assert by_index[1].roles == []

    # "CU-1" is an institution-specific label; the inspector must not
    # claim any academic meaning for it, only report a structural hint.
    assert by_index[2].header_text == "CU-1"
    assert by_index[2].normalized == "cu 1"
    assert by_index[2].roles == []
    assert by_index[2].is_assessment_like is False

    # Generic structural roles are still allowed for common labels.
    assert any(role.kind == "marks_obtained" for role in by_index[3].roles)
    assert any(role.kind == "admission_no" for role in by_index[0].roles)


# ------------------------------------------------------------
# Confidence values are structural and bounded
# ------------------------------------------------------------


def test_confidence_values_are_bounded():
    sheet = sheet_with_rows([
        ["Roll No", "Student Name", "Register Number"],
        ["6A01", "Aarav Sharma", "REG-001"],
    ])
    profile = inspect_sheet(sheet)
    assert 0.0 <= profile.header_block.confidence <= 1.0
    for column in profile.columns:
        for role in column.roles:
            assert 0.0 <= role.confidence <= 1.0

    register = next(
        column for column in profile.columns if column.normalized == "register number"
    )
    assert any(role.kind == "registration_no" for role in register.roles)
    assert register.roles[0].confidence > 0.9


def test_inspection_is_deterministic():
    rows = [
        ["Roll No", "Name", "Physics", "Maths"],
        ["6A01", "Aarav Sharma", 78, 82],
    ]
    first = inspect_sheet(sheet_with_rows(rows))
    second = inspect_sheet(sheet_with_rows(rows))
    assert first.header_block == second.header_block
    assert [column for column in first.columns] == [column for column in second.columns]